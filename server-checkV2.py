import socket
import csv
from concurrent.futures import ThreadPoolExecutor, wait, FIRST_COMPLETED
from pathlib import Path


# ==========================================================
# CONFIG
# ==========================================================

TIMEOUT = 2
MAX_PENDING = 200

SMTP_OUTPUT = "smtp-spaceship.csv"
IMAP_OUTPUT = "imap-spaceship.csv"
POP3_OUTPUT = "pop3-spaceship.csv"


# ==========================================================
# MAIL SERVER PATTERNS
# ==========================================================

SMTP_SERVERS = [
	('smtp1.{}', 25, 0),
	('smtp2.{}', 25, 0),
	('smtp3.{}', 25, 0),
	('smtp4.{}', 587, 1),
	('smtp4.{}', 25, 0),
	('smtp5.{}', 25, 0),
	('smtp6.{}', 25, 0),
	('smtp7.{}', 25, 0),
	('smtp8.{}', 25, 0),
	('smtp9.{}', 25, 0),
	('smtp10.{}', 25, 0),
	('smtp11.{}', 25, 0),
	('smtp12.{}', 25, 0),
	('smtp13.{}', 25, 0),
	('smtp14.{}', 25, 0),
	('smtp16.{}', 25, 0),
	('smtp17.{}', 25, 0),
	('smtp18.{}', 25, 0),
	('smtp20.{}', 25, 0),
	('mail1.{}', 25, 0),
	('mail2.{}', 25, 0),
	('mail3.{}', 25, 0),
	('mail4.{}', 25, 0),
	('mail5.{}', 25, 0),
	('mail6.{}', 25, 0),
	('mail8.{}', 25, 0),
	('mail9.{}', 25, 0),
	('mail10.{}', 25, 0),
	('mail11.{}', 25, 0),
	('mail12.{}', 25, 0),
	('mail13.{}', 25, 0),
	('mail14.{}', 25, 0),
	('mail16.{}', 25, 0),
	('mail17.{}', 25, 0),
	('mail18.{}', 25, 0),
	('mail20.{}', 25, 0),
	('mx1.{}', 25, 0),
	('mx2.{}', 25, 0),
	('mx3.{}', 25, 0),
	('mx4.{}', 25, 0),
	('mx5.{}', 25, 0),
	('mx7.{}', 25, 0),
	('mx8.{}', 25, 0),
	('mx9.{}', 25, 0),
	('mx11.{}', 25, 0),
	('mx12.{}', 25, 0),
	('mx13.{}', 25, 0),
	('mx15.{}', 25, 0),
	('mx16.{}', 25, 0),
	('mx17.{}', 25, 0),
	('mx19.{}', 25, 0),
	('mx20.{}', 25, 0),
	('mx01.{}', 25, 0),
	('mx02.{}', 25, 0),
	('mx03.{}', 25, 0),
	('mx04.{}', 25, 0),
	('mx05.{}', 25, 0),
	('mx06.{}', 25, 0),
	('mx08.{}', 25, 0),
	('mx09.{}', 25, 0),
	('mx14.{}', 25, 0),
	('mx18.{}', 25, 0),
	('smtp.{}', 25, 0),
	('smtp.{}', 465, 1),
	('smtp.{}', 587, 1),
	('smtp.{}', 2525, 0),
	('smtp.{}', 993, 1),
	('smtp.{}', 995, 1),
	('smtp.{}', 110, 0),
	('smtp.{}', 143, 0),
	('smtp.{}', 10025, 0),
	('mail.{}', 25, 0),
	('mail.{}', 465, 1),
	('mail.{}', 587, 1),
	('mail.{}', 2525, 0),
	('mail.{}', 993, 1),
	('mail.{}', 995, 1),
	('mail.{}', 110, 0),
	('mail.{}', 143, 0),
	('mail.{}', 10025, 0),
	('mx.{}', 25, 0),
	('mx.{}', 465, 1),
	('mx.{}', 587, 1),
	('mx.{}', 2525, 0),
	('mx.{}', 10025, 0),
	('email.{}', 25, 0),
	('email.{}', 465, 1),
	('email.{}', 587, 1),
	('email.{}', 2525, 0),
	('relay.{}', 25, 0),
	('relay.{}', 465, 1),
	('relay.{}', 587, 1),
	('relay.{}', 2525, 0),
	('relay.{}', 10025, 0),
	('smtp15.{}', 25, 0),
	('smtp19.{}', 25, 0),
	('mail7.{}', 25, 0),
	('mail19.{}', 25, 0),
	('mx6.{}', 25, 0),
	('mx10.{}', 25, 0),
	('mxs15.{}', 25, 0),
	('mx07.{}', 25, 0),
	('smtp-relay.{}', 25, 0),
	('smtprelay.{}', 25, 0),
	('mail-relay.{}', 25, 0),
	('mailrelay.{}', 25, 0),
	('smtp-mail.{}', 587, 1),
	('smtp-mail.{}', 25, 0),
	('smtpmail.{}', 25, 0),
	('smtp-out.{}', 25, 0),
	('smtpout.{}', 25, 0),
	('mail-out.{}', 25, 0),
	('mailout.{}', 25, 0),
	('smtp-in.{}', 25, 0),
	('smtpin.{}', 25, 0),
	('mail-in.{}', 25, 0),
	('mailin.{}', 25, 0),
	('smtp01.{}', 25, 0),
	('smtp02.{}', 25, 0),
	('mail01.{}', 25, 0),
	('mail02.{}', 25, 0),
	('smtp.secureserver.net.{}', 25, 0),
	('smtp-gw.{}', 25, 0),
	('smtpgw.{}', 25, 0),
	('mail-gw.{}', 25, 0),
	('mailgw.{}', 25, 0),
	('smtp-gateway.{}', 25, 0),
	('smtpgateway.{}', 25, 0),
	('mail-gateway.{}', 25, 0),
	('mailgateway.{}', 25, 0),
	('smtp-server.{}', 25, 0),
	('smtpserver.{}', 25, 0),
	('mail-server.{}', 25, 0),
	('mailserver.{}', 25, 0),
	('smtpauth.{}', 587, 1),
	('smtp-auth.{}', 587, 1),
	('smtpauth.{}', 465, 1),
	('secure-smtp.{}', 465, 1),
	('securesmtp.{}', 465, 1),
	('secure-mail.{}', 465, 1),
	('securemail.{}', 465, 1),
	('smtp-secure.{}', 465, 1),
	('smtpsecure.{}', 465, 1),
	('smtps.{}', 465, 1),
	('smtp-ssl.{}', 465, 1),
	('smtpssl.{}', 465, 1),
	('mail-ssl.{}', 465, 1),
	('mailssl.{}', 465, 1),
	('exchange.{}', 587, 1),
	('exchange.{}', 25, 0),
	('mailhost.{}', 25, 0),
	('mail-host.{}', 25, 0),
	('mailhub.{}', 25, 0),
	('mail-hub.{}', 25, 0),
	('smtprelay.{}', 587, 1),
	('smtp-relay.{}', 587, 1),
	('relay-mail.{}', 25, 0),
	('relaymail.{}', 25, 0),
	('relay-smtp.{}', 25, 0),
	('relaysmtp.{}', 25, 0),
	('mta.{}', 25, 0),
	('mta1.{}', 25, 0),
	('mta2.{}', 25, 0),
	('sendmail.{}', 25, 0),
	('postfix.{}', 25, 0),
	('exim.{}', 25, 0),
	('zimbra.{}', 25, 0),
	('zimbra.{}', 587, 1),
	('mx-relay.{}', 25, 0),
	('mxrelay.{}', 25, 0),
	('mx-gw.{}', 25, 0),
	('mxgw.{}', 25, 0),
	('gateway.{}', 25, 0),
	('gw.{}', 25, 0),
	('smtp-filter.{}', 25, 0),
	('smtpfilter.{}', 25, 0),
	('mail-filter.{}', 25, 0),
	('mailfilter.{}', 25, 0),
	('smtp-proxy.{}', 25, 0),
	('smtpproxy.{}', 25, 0),
	('mail-proxy.{}', 25, 0),
	('mailproxy.{}', 25, 0),
	('smtp-scan.{}', 25, 0),
	('mail-scan.{}', 25, 0),
	('mailscanner.{}', 25, 0),
	('smtp-out1.{}', 25, 0),
	('smtp-out2.{}', 25, 0),
	('smtp-in1.{}', 25, 0),
	('smtp-in2.{}', 25, 0),
	('mail-out1.{}', 25, 0),
	('mail-out2.{}', 25, 0),
	('mail-in1.{}', 25, 0),
	('mail-in2.{}', 25, 0),
	('outbound.{}', 25, 0),
	('outbound-mail.{}', 25, 0),
	('inbound.{}', 25, 0),
	('inbound-mail.{}', 25, 0),
	('smtp-outbound.{}', 25, 0),
	('smtp-inbound.{}', 25, 0),
	('mail-outbound.{}', 25, 0),
	('mail-inbound.{}', 25, 0),
	('send.{}', 25, 0),
	('send-mail.{}', 25, 0),
	('smtp-send.{}', 25, 0),
	('mail-send.{}', 25, 0),
	('out.{}', 25, 0),
	('out-mail.{}', 25, 0),
	('out-smtp.{}', 25, 0),
	('in-mail.{}', 25, 0),
	('in-smtp.{}', 25, 0),
	('out-mail.{}', 587, 1),
	('out-smtp.{}', 587, 1),
	('alt1.aspmx.l.google.{}', 25, 0),
	('alt2.aspmx.l.google.{}', 25, 0),
	('aspmx.l.google.{}', 25, 0),
	('smtp.gmail.{}', 587, 1),
	('smtp.office365.{}', 587, 1),
	('smtp.live.{}', 587, 1),
	('smtp.mail.yahoo.{}', 465, 1),
	('smtp.mail.yahoo.{}', 587, 1),
	('smtp.aol.{}', 465, 1),
	('smtp.zoho.{}', 465, 1),
	('smtp.zoho.{}', 587, 1),
	('smtp.icloud.{}', 587, 1),
	('smtp.mail.me.{}', 587, 1),
	('exchange.{}', 465, 1),
	('exchange-mail.{}', 25, 0),
	('exchange-smtp.{}', 25, 0),
	('exchange-relay.{}', 25, 0),
	('exch.{}', 25, 0),
	('exch1.{}', 25, 0),
	('exch2.{}', 25, 0),
	('smtp-mx.{}', 25, 0),
	('smtpmx.{}', 25, 0),
	('mail-mx.{}', 25, 0),
	('mailmx.{}', 25, 0),
	('mx-mail.{}', 25, 0),
	('mxmail.{}', 25, 0),
	('smtp-gw1.{}', 25, 0),
	('smtp-gw2.{}', 25, 0),
	('mail-gw1.{}', 25, 0),
	('mail-gw2.{}', 25, 0),
	('mx-gw1.{}', 25, 0),
	('mx-gw2.{}', 25, 0),
	('relay1.{}', 25, 0),
	('relay2.{}', 25, 0),
	('relay3.{}', 25, 0),
	('relay01.{}', 25, 0),
	('relay02.{}', 25, 0),
	('relay-out.{}', 25, 0),
	('relayout.{}', 25, 0),
	('relay-in.{}', 25, 0),
	('relayin.{}', 25, 0),
	('smtp-relay1.{}', 25, 0),
	('smtp-relay2.{}', 25, 0),
	('mail-relay1.{}', 25, 0),
	('mail-relay2.{}', 25, 0),
	('backup-mail.{}', 25, 0),
	('backup-smtp.{}', 25, 0),
	('backup-mx.{}', 25, 0),
	('backuprelay.{}', 25, 0),
	('backup-relay.{}', 25, 0),
	('secondary-mail.{}', 25, 0),
	('secondary-mx.{}', 25, 0),
	('secondary-smtp.{}', 25, 0),
	('alt-mail.{}', 25, 0),
	('alt-smtp.{}', 25, 0),
	('alt-mx.{}', 25, 0),
	('fallback-mail.{}', 25, 0),
	('fallback-smtp.{}', 25, 0),
	('fallback-mx.{}', 25, 0),
	('failover-mail.{}', 25, 0),
	('failover-smtp.{}', 25, 0),
	('failover-mx.{}', 25, 0),
	('smtp-alt.{}', 25, 0),
	('smtp-alt.{}', 587, 1),
	('smtpalt.{}', 25, 0),
	('mail-alt.{}', 25, 0),
	('mailalt.{}', 25, 0),
	('mx-backup.{}', 25, 0),
	('mxbackup.{}', 25, 0),
	('smtp-eu.{}', 25, 0),
	('smtp-us.{}', 25, 0),
	('smtp-asia.{}', 25, 0),
	('smtp-eu1.{}', 25, 0),
	('smtp-eu2.{}', 25, 0),
	('smtp-us1.{}', 25, 0),
	('smtp-us2.{}', 25, 0),
	('mail-eu.{}', 25, 0),
	('mail-us.{}', 25, 0),
	('mail-asia.{}', 25, 0),
	('mx-eu.{}', 25, 0),
	('mx-us.{}', 25, 0),
	('mx-asia.{}', 25, 0),
	('smtp-east.{}', 25, 0),
	('smtp-west.{}', 25, 0),
	('mail-east.{}', 25, 0),
	('mail-west.{}', 25, 0),
	('mx-east.{}', 25, 0),
	('mx-west.{}', 25, 0),
	('smtp-uk.{}', 25, 0),
	('smtp-de.{}', 25, 0),
	('smtp-fr.{}', 25, 0),
	('smtp-jp.{}', 25, 0),
	('smtp-au.{}', 25, 0),
	('smtp-ca.{}', 25, 0),
	('mail-uk.{}', 25, 0),
	('mail-de.{}', 25, 0),
	('mail-fr.{}', 25, 0),
	('mail-jp.{}', 25, 0),
	('mail-au.{}', 25, 0),
	('barracuda.{}', 25, 0),
	('mimecast.{}', 25, 0),
	('proofpoint.{}', 25, 0),
	('pphosted.{}', 25, 0),
	('messagelabs.{}', 25, 0),
	('spamfilter.{}', 25, 0),
	('spam-filter.{}', 25, 0),
	('antispam.{}', 25, 0),
	('anti-spam.{}', 25, 0),
	('emailsecurity.{}', 25, 0),
	('email-security.{}', 25, 0),
	('securitygateway.{}', 25, 0),
	('security-gateway.{}', 25, 0),
	('contentfilter.{}', 25, 0),
	('content-filter.{}', 25, 0),
	('ironport.{}', 25, 0),
	('fortimail.{}', 25, 0),
	('smtp-test.{}', 25, 0),
	('smtp-dev.{}', 25, 0),
	('smtp-staging.{}', 25, 0),
	('smtp-prod.{}', 587, 1),
	('mail-test.{}', 25, 0),
	('mail-dev.{}', 25, 0),
	('mail-staging.{}', 25, 0),
	('mail-prod.{}', 25, 0),
	('dev-mail.{}', 25, 0),
	('dev-smtp.{}', 25, 0),
	('test-mail.{}', 25, 0),
	('test-smtp.{}', 25, 0),
	('stage-mail.{}', 25, 0),
	('stage-smtp.{}', 25, 0),
	('prod-mail.{}', 25, 0),
	('prod-smtp.{}', 587, 1),
	('relay4.{}', 25, 0),
	('relay5.{}', 25, 0),
	('relay6.{}', 25, 0),
	('relay7.{}', 25, 0),
	('relay8.{}', 25, 0),
	('relay9.{}', 25, 0),
	('relay10.{}', 25, 0),
	('smtp.sendgrid.{}', 587, 1),
	('smtp.mailgun.{}', 587, 1),
	('smtp.mandrill.{}', 587, 1),
	('smtp.postmark.{}', 587, 1),
	('smtp.sparkpost.{}', 587, 1),
	('smtp.amazonses.{}', 587, 1),
	('email-smtp.{}', 587, 1),
	('ses.{}', 587, 1),
	('sendgrid.{}', 25, 0),
	('mailgun.{}', 25, 0),
	('mailgun.{}', 587, 1),
	('mandrill.{}', 587, 1),
	('postmark.{}', 587, 1),
	('sparkpost.{}', 587, 1),
	('amazonses.{}', 587, 1),
	('cpanel.{}', 25, 0),
	('cpanel.{}', 465, 1),
	('cpanel.{}', 587, 1),
	('plesk.{}', 25, 0),
	('plesk.{}', 587, 1),
	('whm.{}', 25, 0),
	('directadmin.{}', 25, 0),
	('ispconfig.{}', 25, 0),
	('smtp-tls.{}', 587, 1),
	('smtptls.{}', 587, 1),
	('mail-tls.{}', 587, 1),
	('mailtls.{}', 587, 1),
	('tls-smtp.{}', 587, 1),
	('tlssmtp.{}', 587, 1),
	('tls-mail.{}', 587, 1),
	('tlsmail.{}', 587, 1),
	('ssl-smtp.{}', 465, 1),
	('sslsmtp.{}', 465, 1),
	('ssl-mail.{}', 465, 1),
	('sslmail.{}', 465, 1),
	('starttls.{}', 587, 1),
	('smtps-relay.{}', 465, 1),
	('smtps-mail.{}', 465, 1),
	('smtps-relay.{}', 587, 1),
	('smtps-mail.{}', 587, 1),
	('smtp-client.{}', 587, 1),
	('smtpclient.{}', 587, 1),
	('mail-client.{}', 25, 0),
	('mailclient.{}', 25, 0),
	('mailer.{}', 25, 0),
	('mailer1.{}', 25, 0),
	('mailer2.{}', 25, 0),
	('post.{}', 25, 0),
	('post-mail.{}', 25, 0),
	('postmail.{}', 25, 0),
	('post-smtp.{}', 25, 0),
	('postsmtp.{}', 25, 0),
	('postfix-mail.{}', 25, 0),
	('postfix-smtp.{}', 25, 0),
	('postbox.{}', 25, 0),
	('mailbox.{}', 25, 0),
	('mail-box.{}', 25, 0),
	('inbox.{}', 25, 0),
	('outbox.{}', 25, 0),
	('postmaster.{}', 25, 0),
	('mailer-daemon.{}', 25, 0),
	('maildaemon.{}', 25, 0),
	('smtp.corp.{}', 25, 0),
	('mail.corp.{}', 25, 0),
	('corp-mail.{}', 25, 0),
	('corpmail.{}', 25, 0),
	('corp-smtp.{}', 25, 0),
	('corpsmtp.{}', 25, 0),
	('smtp.internal.{}', 25, 0),
	('mail.internal.{}', 25, 0),
	('internal-mail.{}', 25, 0),
	('internalmail.{}', 25, 0),
	('internal-smtp.{}', 25, 0),
	('internalsmtp.{}', 25, 0),
	('smtp.local.{}', 25, 0),
	('mail.local.{}', 25, 0),
	('smtp.dmz.{}', 25, 0),
	('mail.dmz.{}', 25, 0),
	('dmz-mail.{}', 25, 0),
	('dmz-smtp.{}', 25, 0),
	('edge-mail.{}', 25, 0),
	('edge-smtp.{}', 25, 0),
	('edge-mx.{}', 25, 0),
	('perimeter-mail.{}', 25, 0),
	('perimeter-smtp.{}', 25, 0),
	('border-mail.{}', 25, 0),
	('border-smtp.{}', 25, 0),
	('front-mail.{}', 25, 0),
	('front-smtp.{}', 25, 0),
	('frontend-mail.{}', 25, 0),
	('frontend-smtp.{}', 25, 0),
	('backend-mail.{}', 25, 0),
	('backend-smtp.{}', 25, 0),
	('smtp-aws.{}', 25, 0),
	('smtp-azure.{}', 25, 0),
	('smtp-gcp.{}', 25, 0),
	('mail-aws.{}', 25, 0),
	('mail-azure.{}', 25, 0),
	('mail-gcp.{}', 25, 0),
	('cloud-mail.{}', 25, 0),
	('cloud-smtp.{}', 587, 1),
	('cloud-email.{}', 25, 0),
	('hosted-mail.{}', 25, 0),
	('hosted-smtp.{}', 587, 1),
	('managed-mail.{}', 25, 0),
	('managed-smtp.{}', 587, 1),
	('smtp-hosted.{}', 587, 1),
	('mail-hosted.{}', 25, 0),
	('mail-hosting.{}', 25, 0),
	('smtp-hosting.{}', 25, 0),
	('dc1-mail.{}', 25, 0),
	('dc2-mail.{}', 25, 0),
	('dc1-smtp.{}', 25, 0),
	('dc2-smtp.{}', 25, 0),
	('dc-mail.{}', 25, 0),
	('dc-smtp.{}', 25, 0),
	('datacenter-mail.{}', 25, 0),
	('datacenter-smtp.{}', 25, 0),
	('smtp-az1.{}', 25, 0),
	('smtp-az2.{}', 25, 0),
	('mail-az1.{}', 25, 0),
	('mail-az2.{}', 25, 0),
	('smtp-nyc.{}', 25, 0),
	('smtp-lon.{}', 25, 0),
	('smtp-tok.{}', 25, 0),
	('smtp-par.{}', 25, 0),
	('smtp-fra.{}', 25, 0),
	('smtp-sfo.{}', 25, 0),
	('smtp-ams.{}', 25, 0),
	('mailservice.{}', 25, 0),
	('mail-service.{}', 25, 0),
	('mailsvc.{}', 25, 0),
	('mail-svc.{}', 25, 0),
	('mailapp.{}', 25, 0),
	('mail-app.{}', 25, 0),
	('mailagent.{}', 25, 0),
	('mail-agent.{}', 25, 0),
	('mailsrv.{}', 25, 0),
	('mailsrv1.{}', 25, 0),
	('mailsrv2.{}', 25, 0),
	('smtpsrv.{}', 25, 0),
	('smtpsrv1.{}', 25, 0),
	('smtpsrv2.{}', 25, 0),
	('emailserver.{}', 25, 0),
	('emailserver.{}', 587, 1),
	('email-server.{}', 25, 0),
	('emailhost.{}', 25, 0),
	('email-host.{}', 25, 0),
	('auth-mail.{}', 25, 0),
	('authmail.{}', 25, 0),
	('auth-smtp.{}', 587, 1),
	('authsmtp.{}', 587, 1),
	('login-mail.{}', 25, 0),
	('login-smtp.{}', 587, 1),
	('authenticated-mail.{}', 25, 0),
	('authenticated-smtp.{}', 587, 1),
	('mail-api.{}', 25, 0),
	('mailapi.{}', 25, 0),
	('smtp-api.{}', 587, 1),
	('smtpapi.{}', 587, 1),
	('email-api.{}', 25, 0),
	('emailapi.{}', 25, 0),
	('api-mail.{}', 25, 0),
	('apimail.{}', 25, 0),
	('api-smtp.{}', 587, 1),
	('apis.{}', 587, 1),
	('api-email.{}', 587, 1),
	('apiemail.{}', 587, 1),
	('notification.{}', 587, 1),
	('notifications.{}', 587, 1),
	('notify.{}', 587, 1),
	('alert.{}', 587, 1),
	('alerts.{}', 587, 1),
	('noreply.{}', 587, 1),
	('no-reply.{}', 587, 1),
	('donotreply.{}', 587, 1),
	('system-mail.{}', 25, 0),
	('systemmail.{}', 25, 0),
	('system-smtp.{}', 587, 1),
	('systemsmtp.{}', 587, 1),
	('cron-mail.{}', 25, 0),
	('cronmail.{}', 25, 0),
	('auto-mail.{}', 25, 0),
	('automail.{}', 25, 0),
	('lb-mail.{}', 25, 0),
	('lbmail.{}', 25, 0),
	('lb-smtp.{}', 25, 0),
	('lb1-mail.{}', 25, 0),
	('lb2-mail.{}', 25, 0),
	('cluster-mail.{}', 25, 0),
	('clustermail.{}', 25, 0),
	('cluster-smtp.{}', 25, 0),
	('node-mail.{}', 25, 0),
	('node1-mail.{}', 25, 0),
	('node2-mail.{}', 25, 0),
	('pool-mail.{}', 25, 0),
	('pool-smtp.{}', 25, 0),
	('shard-mail.{}', 25, 0),
	('shard1-mail.{}', 25, 0),
	('shard2-mail.{}', 25, 0),
	('docker-mail.{}', 25, 0),
	('docker-smtp.{}', 25, 0),
	('k8s-mail.{}', 25, 0),
	('k8s-smtp.{}', 25, 0),
	('pod-mail.{}', 25, 0),
	('pod-smtp.{}', 25, 0),
	('container-mail.{}', 25, 0),
	('container-smtp.{}', 25, 0),
	('support-mail.{}', 25, 0),
	('supportmail.{}', 25, 0),
	('helpdesk-mail.{}', 25, 0),
	('helpdeskmail.{}', 25, 0),
	('sales-mail.{}', 25, 0),
	('salesmail.{}', 25, 0),
	('billing-mail.{}', 25, 0),
	('billingmail.{}', 25, 0),
	('crm-mail.{}', 25, 0),
	('crmmail.{}', 25, 0),
	('hr-mail.{}', 25, 0),
	('hrmail.{}', 25, 0),
	('info-mail.{}', 25, 0),
	('infomail.{}', 25, 0),
	('contact-mail.{}', 25, 0),
	('contactmail.{}', 25, 0),
	('invoice-mail.{}', 25, 0),
	('invoicemail.{}', 25, 0),
	('marketing-mail.{}', 25, 0),
	('marketingsmtp.{}', 587, 1),
	('marketing-smtp.{}', 587, 1),
	('campaign-mail.{}', 25, 0),
	('campaign-smtp.{}', 587, 1),
	('bulk-mail.{}', 25, 0),
	('bulkmail.{}', 25, 0),
	('bulk-smtp.{}', 587, 1),
	('bulksmtp.{}', 587, 1),
	('mass-mail.{}', 25, 0),
	('massmail.{}', 25, 0),
	('newsletter.{}', 25, 0),
	('newsletter-mail.{}', 25, 0),
	('transactional.{}', 587, 1),
	('transactional-mail.{}', 25, 0),
	('transactional-smtp.{}', 587, 1),
	('verify-mail.{}', 587, 1),
	('verify-smtp.{}', 587, 1),
	('verify.{}', 587, 1),
	('confirm-mail.{}', 587, 1),
	('confirm-smtp.{}', 587, 1),
	('confirm.{}', 587, 1),
	('activation-mail.{}', 587, 1),
	('activation-smtp.{}', 587, 1),
	('reset-mail.{}', 587, 1),
	('reset-smtp.{}', 587, 1),
	('password-mail.{}', 587, 1),
	('password-smtp.{}', 587, 1),
	('recover-mail.{}', 587, 1),
	('recovery-smtp.{}', 587, 1),
	('welcome-mail.{}', 587, 1),
	('welcome-smtp.{}', 587, 1),
	('correo.{}', 25, 0),
	('correo.{}', 587, 1),
	('correio.{}', 25, 0),
	('poczta.{}', 25, 0),
	('posta.{}', 25, 0),
	('brief.{}', 25, 0),
	('epost.{}', 25, 0),
	('poste.{}', 25, 0),
	('courriel.{}', 25, 0),
	('smtp.mail.{}', 25, 0),
	('smtp.mail.{}', 587, 1),
	('mail.smtp.{}', 25, 0),
	('mx.mail.{}', 25, 0),
	('smtp.mx.{}', 25, 0),
	('mail.mx.{}', 25, 0),
	('smtp.relay.{}', 25, 0),
	('mail.relay.{}', 25, 0),
	('relay.mail.{}', 25, 0),
	('relay.smtp.{}', 25, 0),
	('mx.smtp.{}', 25, 0),
	('mx.relay.{}', 25, 0),
	('smtp.email.{}', 25, 0),
	('smtp.email.{}', 587, 1),
	('mail.email.{}', 25, 0),
	('email.mail.{}', 25, 0),
	('email.smtp.{}', 25, 0),
	('smtp.post.{}', 25, 0),
	('mail.post.{}', 25, 0),
	('mx.email.{}', 25, 0),
	('smtp1.mail.{}', 25, 0),
	('smtp2.mail.{}', 25, 0),
	('mx1.mail.{}', 25, 0),
	('mx2.mail.{}', 25, 0),
	('relay1.mail.{}', 25, 0),
	('relay2.mail.{}', 25, 0),
	('mail1.mail.{}', 25, 0),
	('mail2.mail.{}', 25, 0),
	('esmtp.{}', 25, 0),
	('esmtp.{}', 587, 1),
	('esmtp-relay.{}', 25, 0),
	('smtpd.{}', 25, 0),
	('smtpd.{}', 587, 1),
	('qmail.{}', 25, 0),
	('qmail-smtp.{}', 25, 0),
	('exim-smtp.{}', 25, 0),
	('exim4.{}', 25, 0),
	('sendmail-server.{}', 25, 0),
	('opensmtpd.{}', 25, 0),
	('haraka.{}', 25, 0),
	('dovecot.{}', 25, 0),
	('dovecot-mail.{}', 25, 0),
	('cyrus.{}', 25, 0),
	('cyrus-mail.{}', 25, 0),
	('maildrop.{}', 25, 0),
	('mail-drop.{}', 25, 0),
	('maildrop1.{}', 25, 0),
	('maildrop2.{}', 25, 0),
	('mailq.{}', 25, 0),
	('mail-queue.{}', 25, 0),
	('mailqueue.{}', 25, 0),
	('queue-mail.{}', 25, 0),
	('queuemail.{}', 25, 0),
	('forward-mail.{}', 25, 0),
	('forwardmail.{}', 25, 0),
	('forwarder-mail.{}', 25, 0),
	('forwardersmtp.{}', 25, 0),
	('bounce-mail.{}', 25, 0),
	('bouncemail.{}', 25, 0),
	('bounce-smtp.{}', 587, 1),
	('bouncesmtp.{}', 587, 1),
	('redirect-mail.{}', 25, 0),
	('redirectmail.{}', 25, 0),
	('dr-mail.{}', 25, 0),
	('dr-smtp.{}', 25, 0),
	('dr-mx.{}', 25, 0),
	('dr-relay.{}', 25, 0),
	('disaster-mail.{}', 25, 0),
	('disaster-smtp.{}', 25, 0),
	('contingency-mail.{}', 25, 0),
	('contingency-smtp.{}', 25, 0),
	('standby-mail.{}', 25, 0),
	('standby-smtp.{}', 25, 0),
	('standby-mx.{}', 25, 0),
	('mirror-mail.{}', 25, 0),
	('mirror-smtp.{}', 25, 0),
	('mirror-mx.{}', 25, 0),
	('replica-mail.{}', 25, 0),
	('replica-smtp.{}', 25, 0),
	('replica-mx.{}', 25, 0),
	('smtp-alt.{}', 2525, 0),
	('mail-alt.{}', 2525, 0),
	('smtp-alternate.{}', 2525, 0),
	('mail-alternate.{}', 2525, 0),
	('alt-smtp.{}', 2525, 0),
	('alt-port.{}', 2525, 0),
	('smtp-relay-gw.{}', 25, 0),
	('mail-relay-gw.{}', 25, 0),
	('relay-gw-mail.{}', 25, 0),
	('smtp-gw-relay.{}', 25, 0),
	('mail-gw-relay.{}', 25, 0),
	('gw-relay-mail.{}', 25, 0),
	('smtp-gateway-1.{}', 25, 0),
	('smtp-gateway-2.{}', 25, 0),
	('mail-gateway-1.{}', 25, 0),
	('mail-gateway-2.{}', 25, 0),
	('smtpgw1.{}', 25, 0),
	('smtpgw2.{}', 25, 0),
	('mailgw1.{}', 25, 0),
	('mailgw2.{}', 25, 0),
	('mxgw1.{}', 25, 0),
	('mxgw2.{}', 25, 0),
	('relaygw1.{}', 25, 0),
	('relaygw2.{}', 25, 0),
	('mail-mta.{}', 25, 0),
	('mailmta.{}', 25, 0),
	('smtp-mta.{}', 25, 0),
	('smtpmta.{}', 25, 0),
	('mta-mail.{}', 25, 0),
	('mtasmtp.{}', 25, 0),
	('mta-relay.{}', 25, 0),
	('mtarelay.{}', 25, 0),
	('mta-in.{}', 25, 0),
	('mta-out.{}', 25, 0),
	('mta1-mail.{}', 25, 0),
	('mta2-mail.{}', 25, 0),
	('mail.postoffice.{}', 25, 0),
	('postoffice.{}', 25, 0),
	('post-office.{}', 25, 0),
	('postbox-server.{}', 25, 0),
	('courier-mail.{}', 25, 0),
	('courier-smtp.{}', 25, 0),
	('sendmail-relay.{}', 25, 0),
	('postfix-relay.{}', 25, 0),
	('exim-relay.{}', 25, 0),
	('zimbra-relay.{}', 25, 0),
	('zimbra-smtp.{}', 25, 0),
	('zimbra-mail.{}', 25, 0),
	('mailcow.{}', 25, 0),
	('mailcow-smtp.{}', 25, 0),
	('mailu.{}', 25, 0),
	('mailu-smtp.{}', 25, 0),
	('iredmail.{}', 25, 0),
	('iredmail-smtp.{}', 25, 0),
	('modoboa.{}', 25, 0),
	('docker-mailserver.{}', 25, 0),
	('sogo.{}', 25, 0),
	('sogo-mail.{}', 25, 0),
	('app-mail.{}', 25, 0),
	('appmail.{}', 25, 0),
	('app-smtp.{}', 587, 1),
	('appsmtp.{}', 587, 1),
	('service-mail.{}', 25, 0),
	('servicemail.{}', 25, 0),
	('service-smtp.{}', 587, 1),
	('servicesmtp.{}', 587, 1),
	('svc-mail.{}', 25, 0),
	('svcmail.{}', 25, 0),
	('svc-smtp.{}', 587, 1),
	('svcsmtp.{}', 587, 1),
	('smtp-v4.{}', 25, 0),
	('smtp-v6.{}', 25, 0),
	('mail-v4.{}', 25, 0),
	('mail-v6.{}', 25, 0),
	('mx-v4.{}', 25, 0),
	('mx-v6.{}', 25, 0),
	('smtp-a.{}', 25, 0),
	('smtp-b.{}', 25, 0),
	('smtp-c.{}', 25, 0),
	('mail-a.{}', 25, 0),
	('mail-b.{}', 25, 0),
	('mail-c.{}', 25, 0),
	('mx-a.{}', 25, 0),
	('mx-b.{}', 25, 0),
	('mx-c.{}', 25, 0),
	('relay-a.{}', 25, 0),
	('relay-b.{}', 25, 0),
	('smtp1a.{}', 25, 0),
	('smtp1b.{}', 25, 0),
	('smtp2a.{}', 25, 0),
	('smtp2b.{}', 25, 0),
	('mail1a.{}', 25, 0),
	('mail1b.{}', 25, 0),
	('mail2a.{}', 25, 0),
	('mail2b.{}', 25, 0),
	('mx1a.{}', 25, 0),
	('mx1b.{}', 25, 0),
	('mx2a.{}', 25, 0),
	('mx2b.{}', 25, 0),
	('front-mx.{}', 25, 0),
	('edge1-mail.{}', 25, 0),
	('edge2-mail.{}', 25, 0),
	('edge1-smtp.{}', 25, 0),
	('edge2-smtp.{}', 25, 0),
	('vip-mail.{}', 25, 0),
	('vip-smtp.{}', 25, 0),
	('virtual-mail.{}', 25, 0),
	('virtualmail.{}', 25, 0),
	('virtual-smtp.{}', 587, 1),
	('virtualsmtp.{}', 587, 1),
	('virtual-mx.{}', 25, 0),
	('virtualmx.{}', 25, 0),
	('primary-mail.{}', 25, 0),
	('primarymail.{}', 25, 0),
	('primary-smtp.{}', 25, 0),
	('primarysmtp.{}', 25, 0),
	('primary-mx.{}', 25, 0),
	('primarymx.{}', 25, 0),
	('primary-relay.{}', 25, 0),
	('primaryrelay.{}', 25, 0),
	('secondarymail.{}', 25, 0),
	('secondarysmtp.{}', 25, 0),
	('secondary-relay.{}', 25, 0),
	('secondaryrelay.{}', 25, 0),
	('smtp-relay.mail.{}', 25, 0),
	('mail-relay.smtp.{}', 25, 0),
	('smtp-relay.mx.{}', 25, 0),
	('mail-relay.mx.{}', 25, 0),
	('smtp-relay.email.{}', 25, 0),
	('mail-relay.email.{}', 25, 0),
	('relay.mail.smtp.{}', 25, 0),
	('smtp.mail.relay.{}', 25, 0),
	('mail.smtp.mx.{}', 25, 0),
	('mx.mail.smtp.{}', 25, 0),
	('smtp.mx.mail.{}', 25, 0),
	('mail.mx.smtp.{}', 25, 0),
	('mailserver.{}', 587, 1),
	('mailserver.{}', 465, 1),
	('smtpserver.{}', 587, 1),
	('smtpserver.{}', 465, 1),
	('emailserver.{}', 465, 1),
	('relayserver.{}', 587, 1),
	('mxserver.{}', 587, 1),
	('gateway-smtp.{}', 587, 1),
	('gateway-mail.{}', 25, 0),
	('gateway-relay.{}', 25, 0),
	('gateway-mx.{}', 25, 0),
	('smtp-corps.{}', 25, 0),
	('smtp-enterprise.{}', 25, 0),
	('smtp-enterprise.{}', 587, 1),
	('mail-enterprise.{}', 25, 0),
	('ent-mail.{}', 25, 0),
	('ent-smtp.{}', 25, 0),
	('ent-mx.{}', 25, 0),
	('hq-mail.{}', 25, 0),
	('hq-smtp.{}', 25, 0),
	('hq-mx.{}', 25, 0),
	('office-mail.{}', 25, 0),
	('office-smtp.{}', 25, 0),
	('office1-mail.{}', 25, 0),
	('office2-mail.{}', 25, 0),
	('main-mail.{}', 25, 0),
	('mainmail.{}', 25, 0),
	('main-smtp.{}', 25, 0),
	('mainsmtp.{}', 25, 0),
	('main-mx.{}', 25, 0),
	('mainmx.{}', 25, 0),
	('main-relay.{}', 25, 0),
	('mainrelay.{}', 25, 0),
	('core-mail.{}', 25, 0),
	('coremail.{}', 25, 0),
	('core-smtp.{}', 25, 0),
	('coresmtp.{}', 25, 0),
	('core-mx.{}', 25, 0),
	('coremx.{}', 25, 0),
	('core-relay.{}', 25, 0),
	('corerelay.{}', 25, 0),
	('mailhosting.{}', 25, 0),
	('smtphosting.{}', 25, 0),
	('email-hosting.{}', 25, 0),
	('emailhosting.{}', 25, 0),
	('mx-hosting.{}', 25, 0),
	('mxhosting.{}', 25, 0),
	('mailhosted.{}', 25, 0),
	('smtphosted.{}', 587, 1),
	('mx-hosted.{}', 25, 0),
	('mxhosted.{}', 25, 0),
	('email-hosted.{}', 25, 0),
	('emailhosted.{}', 25, 0),
	('relay-hosted.{}', 25, 0),
	('relayhosted.{}', 25, 0),
	('smart-host.{}', 25, 0),
	('smarthost.{}', 25, 0),
	('smart-relay.{}', 25, 0),
	('smartrelay.{}', 25, 0),
	('open-relay.{}', 25, 0),
	('openrelay.{}', 25, 0),
	('null-client.{}', 25, 0),
	('nullclient.{}', 25, 0),
	('null-mail.{}', 25, 0),
	('satellite-mail.{}', 25, 0),
	('satellite-smtp.{}', 25, 0),
	('local-delivery.{}', 25, 0),
	('localdelivery.{}', 25, 0),
	('direct-delivery.{}', 25, 0),
	('directdelivery.{}', 25, 0),
	('secure-relay.{}', 465, 1),
	('secure-mx.{}', 465, 1),
	('secure-email.{}', 465, 1),
	('secureemail.{}', 465, 1),
	('encrypted-mail.{}', 465, 1),
	('encryptedmail.{}', 465, 1),
	('encrypt-smtp.{}', 465, 1),
	('encryptsmtp.{}', 465, 1),
	('mail15.{}', 25, 0),
	('correo1.{}', 25, 0),
	('correo2.{}', 25, 0),
	('correo-relay.{}', 25, 0),
	('correorelay.{}', 25, 0),
	('correio.{}', 587, 1),
	('posta1.{}', 25, 0),
	('briefkasten.{}', 25, 0),
	('nachricht.{}', 25, 0),
	('nachrichten.{}', 25, 0),
	('maila.{}', 25, 0),
	('brev.{}', 25, 0),
	('pochta.{}', 25, 0),
	('youjian.{}', 25, 0),
	('mail.xx.{}', 25, 0),
	('mel.{}', 25, 0),
	('postkasse.{}', 25, 0),
	('e-brief.{}', 25, 0),
	('brievenbus.{}', 25, 0),
	('postbus.{}', 25, 0),
	('kotak.{}', 25, 0),
	('surat.{}', 25, 0),
	('posta-elektronika.{}', 25, 0),
	('poczta1.{}', 25, 0),
	('mejl.{}', 25, 0),
	('ephone.{}', 25, 0),
	('epost.{}', 587, 1),
	('backup.{}', 25, 0),
	('backup1.{}', 25, 0),
	('backup2.{}', 25, 0),
	('backup01.{}', 25, 0),
	('backup02.{}', 25, 0),
	('bak.{}', 25, 0),
	('bak1.{}', 25, 0),
	('bak2.{}', 25, 0),
	('secondary.{}', 25, 0),
	('secondary1.{}', 25, 0),
	('secondary2.{}', 25, 0),
	('alt.{}', 25, 0),
	('alt1.{}', 25, 0),
	('alt2.{}', 25, 0),
	('alt1.mail.{}', 25, 0),
	('alt2.mail.{}', 25, 0),
	('spare-mail.{}', 25, 0),
	('sparemail.{}', 25, 0),
	('spare-smtp.{}', 25, 0),
	('sparesmtp.{}', 25, 0),
	('mx.mail1.{}', 25, 0),
	('mx.mail2.{}', 25, 0),
	('mx.email.{}', 587, 1),
	('mx.post.{}', 25, 0),
	('mx.post.{}', 587, 1),
	('mx1.email.{}', 25, 0),
	('mx2.email.{}', 25, 0),
	('mx-backup1.{}', 25, 0),
	('mx-backup2.{}', 25, 0),
	('mx-backup.{}', 465, 1),
	('mxbackup1.{}', 25, 0),
	('mxbackup2.{}', 25, 0),
	('mx-primary.{}', 25, 0),
	('mx-primary.{}', 587, 1),
	('mxprimary.{}', 25, 0),
	('mx-secondary.{}', 25, 0),
	('mxsecondary.{}', 25, 0),
	('mx-failover.{}', 25, 0),
	('mxfailover.{}', 25, 0),
	('mx-fallback.{}', 25, 0),
	('mxfallback.{}', 25, 0),
	('lb.{}', 25, 0),
	('lb1.{}', 25, 0),
	('lb2.{}', 25, 0),
	('lb1.mail.{}', 25, 0),
	('lb2.mail.{}', 25, 0),
	('cluster.{}', 25, 0),
	('cluster1.{}', 25, 0),
	('cluster2.{}', 25, 0),
	('node1.{}', 25, 0),
	('node2.{}', 25, 0),
	('node3.{}', 25, 0),
	('node1.mail.{}', 25, 0),
	('node2.mail.{}', 25, 0),
	('pool.{}', 25, 0),
	('pool1.{}', 25, 0),
	('pool2.{}', 25, 0),
	('pool1.mail.{}', 25, 0),
	('pool2.mail.{}', 25, 0),
	('shard.{}', 25, 0),
	('shard1.{}', 25, 0),
	('shard2.{}', 25, 0),
	('shard1.mail.{}', 25, 0),
	('shard2.mail.{}', 25, 0),
	('docker-mail.{}', 587, 1),
	('docker-smtp.{}', 587, 1),
	('k8s-mail.{}', 587, 1),
	('k8s-smtp.{}', 587, 1),
	('kubernetes-mail.{}', 25, 0),
	('kubernetes-smtp.{}', 25, 0),
	('mail-container.{}', 25, 0),
	('smtp-container.{}', 25, 0),
	('mail-pod.{}', 25, 0),
	('smtp-pod.{}', 25, 0),
	('email-smtp.{}', 25, 0),
	('ses.{}', 25, 0),
	('aws-ses.{}', 587, 1),
	('aws-mail.{}', 25, 0),
	('aws-smtp.{}', 587, 1),
	('ses-mail.{}', 25, 0),
	('ses-smtp.{}', 587, 1),
	('azure-mail.{}', 25, 0),
	('azure-smtp.{}', 587, 1),
	('gcp-mail.{}', 25, 0),
	('gcp-smtp.{}', 587, 1),
	('sendgrid-mail.{}', 25, 0),
	('sendgrid-smtp.{}', 587, 1),
	('mailgun-smtp.{}', 587, 1),
	('smtp-mailgun.{}', 587, 1),
	('smtp-sendgrid.{}', 587, 1),
	('smtp-mandrill.{}', 587, 1),
	('smtp-postmark.{}', 587, 1),
	('smtp-sparkpost.{}', 587, 1),
	('smtp-amazonses.{}', 587, 1),
	('smtp-mailjet.{}', 587, 1),
	('smtp-pepipost.{}', 587, 1),
	('smtp-socketlabs.{}', 587, 1),
	('smtp-sendinblue.{}', 587, 1),
	('smtp-postmarkapp.{}', 587, 1),
	('sendmail-smtp.{}', 25, 0),
	('exim-server.{}', 25, 0),
	('qmail-server.{}', 25, 0),
	('qmail-relay.{}', 25, 0),
	('zimbra-mail.{}', 587, 1),
	('zimbra-smtp.{}', 587, 1),
	('haraka-smtp.{}', 25, 0),
	('opensmtpd-mail.{}', 25, 0),
	('mailenable.{}', 25, 0),
	('mailenable-smtp.{}', 25, 0),
	('smartermail.{}', 25, 0),
	('smartermail.{}', 587, 1),
	('hmailserver.{}', 25, 0),
	('hmail.{}', 25, 0),
	('axigen.{}', 25, 0),
	('axigen-mail.{}', 25, 0),
	('scalix.{}', 25, 0),
	('soGo.{}', 25, 0),
	('barracuda-mail.{}', 25, 0),
	('barracuda-smtp.{}', 25, 0),
	('barracuda-relay.{}', 25, 0),
	('barracuda1.{}', 25, 0),
	('barracuda2.{}', 25, 0),
	('mimecast-mail.{}', 25, 0),
	('mimecast-smtp.{}', 25, 0),
	('mimecast-relay.{}', 25, 0),
	('mimecast1.{}', 25, 0),
	('mimecast2.{}', 25, 0),
	('proofpoint-mail.{}', 25, 0),
	('proofpoint-smtp.{}', 25, 0),
	('proofpoint-relay.{}', 25, 0),
	('pphosted-mail.{}', 25, 0),
	('pphosted-smtp.{}', 25, 0),
	('symantec-mail.{}', 25, 0),
	('symantec-smtp.{}', 25, 0),
	('symantec-relay.{}', 25, 0),
	('messagelabs-smtp.{}', 25, 0),
	('messagelabs-relay.{}', 25, 0),
	('trendmicro-smtp.{}', 25, 0),
	('mcafee-smtp.{}', 25, 0),
	('sophos-mail.{}', 25, 0),
	('sophos-smtp.{}', 25, 0),
	('fortimail.{}', 587, 1),
	('fortimail-smtp.{}', 25, 0),
	('forcepoint.{}', 25, 0),
	('forcepoint-smtp.{}', 25, 0),
	('websense.{}', 25, 0),
	('websense-mail.{}', 25, 0),
	('fireeye.{}', 25, 0),
	('fireeye-mail.{}', 25, 0),
	('trend-email.{}', 25, 0),
	('trendmicro-email.{}', 25, 0),
	('sonicwall-mail.{}', 25, 0),
	('sonicwall-smtp.{}', 25, 0),
	('watchguard-mail.{}', 25, 0),
	('watchguard-smtp.{}', 25, 0),
	('cisco-mail.{}', 25, 0),
	('cisco-smtp.{}', 25, 0),
	('ironport.{}', 587, 1),
	('ironport-mail.{}', 25, 0),
	('ironport-smtp.{}', 25, 0),
	('cisco-esa.{}', 25, 0),
	('esa-mail.{}', 25, 0),
	('smtp-out-relay.{}', 25, 0),
	('smtp-in-relay.{}', 25, 0),
	('smtp-relay-out.{}', 25, 0),
	('smtp-relay-in.{}', 25, 0),
	('mail-out-relay.{}', 25, 0),
	('mail-in-relay.{}', 25, 0),
	('mail-relay-out.{}', 25, 0),
	('mail-relay-in.{}', 25, 0),
	('relay-out-smtp.{}', 25, 0),
	('relay-in-smtp.{}', 25, 0),
	('mx-in-relay.{}', 25, 0),
	('mx-out-relay.{}', 25, 0),
	('email-relay-out.{}', 25, 0),
	('email-relay-in.{}', 25, 0),
	('outbound-mail-relay.{}', 25, 0),
	('inbound-mail-relay.{}', 25, 0),
	('external-mail.{}', 25, 0),
	('external-smtp.{}', 587, 1),
	('external-relay.{}', 25, 0),
	('external-mx.{}', 25, 0),
	('internal-mx.{}', 25, 0),
	('smtp-sin.{}', 25, 0),
	('smtp-syd.{}', 25, 0),
	('smtp-dub.{}', 25, 0),
	('mail-nyc.{}', 25, 0),
	('mail-lon.{}', 25, 0),
	('mail-tok.{}', 25, 0),
	('mail-par.{}', 25, 0),
	('mail-fra.{}', 25, 0),
	('mail-sin.{}', 25, 0),
	('mail-syd.{}', 25, 0),
	('mail-sfo.{}', 25, 0),
	('mx-nyc.{}', 25, 0),
	('mx-lon.{}', 25, 0),
	('mx-tok.{}', 25, 0),
	('mx-par.{}', 25, 0),
	('mx-fra.{}', 25, 0),
	('mx-sin.{}', 25, 0),
	('smtp-az3.{}', 25, 0),
	('smtp-aza.{}', 25, 0),
	('smtp-azb.{}', 25, 0),
	('smtp-azc.{}', 25, 0),
	('mail-az3.{}', 25, 0),
	('mx-az1.{}', 25, 0),
	('mx-az2.{}', 25, 0),
	('mx-az3.{}', 25, 0),
	('relay-az1.{}', 25, 0),
	('relay-az2.{}', 25, 0),
	('smtp21.{}', 25, 0),
	('smtp22.{}', 25, 0),
	('smtp23.{}', 25, 0),
	('smtp24.{}', 25, 0),
	('smtp25.{}', 25, 0),
	('smtp30.{}', 25, 0),
	('smtp40.{}', 25, 0),
	('smtp50.{}', 25, 0),
	('mail30.{}', 25, 0),
	('mail40.{}', 25, 0),
	('mail50.{}', 25, 0),
	('mx30.{}', 25, 0),
	('mx40.{}', 25, 0),
	('mx50.{}', 25, 0),
	('relay11.{}', 25, 0),
	('relay12.{}', 25, 0),
	('relay13.{}', 25, 0),
	('relay14.{}', 25, 0),
	('relay15.{}', 25, 0),
	('relay16.{}', 25, 0),
	('relay17.{}', 25, 0),
	('relay18.{}', 25, 0),
	('relay19.{}', 25, 0),
	('relay20.{}', 25, 0),
	('mail.mail.{}', 25, 0),
	('smtp.smtp.{}', 25, 0),
	('mx.mx.{}', 25, 0),
	('relay.relay.{}', 25, 0),
	('email.email.{}', 25, 0),
	('mail.mx.mail.{}', 25, 0),
	('smtp.mail.mx.{}', 25, 0),
	('mx.smtp.mail.{}', 25, 0),
	('mail-smtp.relay.{}', 25, 0),
	('mail.relay.smtp.{}', 25, 0),
	('smtp.mail-relay.{}', 25, 0),
	('relay-smtp.mail.{}', 25, 0),
	('email.mail.smtp.{}', 25, 0),
	('smtp.email.mail.{}', 25, 0),
	('mail.smtp.email.{}', 25, 0),
	('app1-mail.{}', 25, 0),
	('app2-mail.{}', 25, 0),
	('app1-smtp.{}', 587, 1),
	('app2-smtp.{}', 587, 1),
	('webapp-mail.{}', 25, 0),
	('webapp-smtp.{}', 587, 1),
	('mobile-mail.{}', 25, 0),
	('mobile-smtp.{}', 587, 1),
	('desktop-mail.{}', 25, 0),
	('desktop-smtp.{}', 587, 1),
	('client-mail.{}', 25, 0),
	('client-smtp.{}', 587, 1),
	('server-mail.{}', 25, 0),
	('server-smtp.{}', 25, 0),
	('micro-mail.{}', 25, 0),
	('micro-smtp.{}', 25, 0),
	('microservice-mail.{}', 25, 0),
	('microservice-smtp.{}', 25, 0),
	('monitor-mail.{}', 25, 0),
	('monitor-smtp.{}', 25, 0),
	('mon-mail.{}', 25, 0),
	('mon-smtp.{}', 25, 0),
	('nagios-mail.{}', 25, 0),
	('nagios-smtp.{}', 25, 0),
	('zabbix-mail.{}', 25, 0),
	('zabbix-smtp.{}', 25, 0),
	('prometheus-mail.{}', 25, 0),
	('prometheus-smtp.{}', 25, 0),
	('grafana-mail.{}', 25, 0),
	('grafana-smtp.{}', 25, 0),
	('jenkins-mail.{}', 25, 0),
	('jenkins-smtp.{}', 25, 0),
	('ci-mail.{}', 25, 0),
	('ci-smtp.{}', 25, 0),
	('cd-mail.{}', 25, 0),
	('cd-smtp.{}', 25, 0),
	('build-mail.{}', 25, 0),
	('build-smtp.{}', 25, 0),
	('deploy-mail.{}', 25, 0),
	('deploy-smtp.{}', 25, 0),
	('git-mail.{}', 25, 0),
	('git-smtp.{}', 25, 0),
	('gitlab-mail.{}', 25, 0),
	('gitlab-smtp.{}', 25, 0),
	('github-mail.{}', 25, 0),
	('bitbucket-mail.{}', 25, 0),
	('ticket-mail.{}', 25, 0),
	('ticket-smtp.{}', 25, 0),
	('ticketing-mail.{}', 25, 0),
	('ticketing-smtp.{}', 25, 0),
	('helpdesk-smtp.{}', 25, 0),
	('support-smtp.{}', 587, 1),
	('desk-mail.{}', 25, 0),
	('desk-smtp.{}', 25, 0),
	('otrs-mail.{}', 25, 0),
	('otrs-smtp.{}', 25, 0),
	('zendesk-mail.{}', 25, 0),
	('zendesk-smtp.{}', 25, 0),
	('freshdesk-mail.{}', 25, 0),
	('freshdesk-smtp.{}', 25, 0),
	('jira-mail.{}', 25, 0),
	('jira-smtp.{}', 25, 0),
	('servicenow-mail.{}', 25, 0),
	('servicenow-smtp.{}', 25, 0),
	('shop-mail.{}', 25, 0),
	('shop-smtp.{}', 587, 1),
	('store-mail.{}', 25, 0),
	('store-smtp.{}', 587, 1),
	('ecommerce-mail.{}', 25, 0),
	('ecommerce-smtp.{}', 587, 1),
	('order-mail.{}', 25, 0),
	('order-smtp.{}', 587, 1),
	('orders-mail.{}', 25, 0),
	('orders-smtp.{}', 587, 1),
	('checkout-mail.{}', 25, 0),
	('checkout-smtp.{}', 587, 1),
	('payment-mail.{}', 25, 0),
	('payment-smtp.{}', 587, 1),
	('payments-mail.{}', 25, 0),
	('payments-smtp.{}', 587, 1),
	('cart-mail.{}', 25, 0),
	('cart-smtp.{}', 587, 1),
	('magento-mail.{}', 25, 0),
	('magento-smtp.{}', 587, 1),
	('shopify-mail.{}', 25, 0),
	('shopify-smtp.{}', 587, 1),
	('woocommerce-mail.{}', 25, 0),
	('woocommerce-smtp.{}', 587, 1),
	('mkt-mail.{}', 25, 0),
	('mkt-smtp.{}', 587, 1),
	('campaigns-mail.{}', 25, 0),
	('campaigns-smtp.{}', 587, 1),
	('promo-mail.{}', 25, 0),
	('promo-smtp.{}', 587, 1),
	('promo1-mail.{}', 25, 0),
	('promo2-mail.{}', 25, 0),
	('ads-mail.{}', 25, 0),
	('ads-smtp.{}', 587, 1),
	('hubspot-mail.{}', 25, 0),
	('hubspot-smtp.{}', 587, 1),
	('salesforce-mail.{}', 25, 0),
	('salesforce-smtp.{}', 587, 1),
	('marketo-mail.{}', 25, 0),
	('marketo-smtp.{}', 587, 1),
	('pardot-mail.{}', 25, 0),
	('pardot-smtp.{}', 587, 1),
	('activecampaign-mail.{}', 25, 0),
	('activecampaign-smtp.{}', 587, 1),
	('getresponse-mail.{}', 25, 0),
	('getresponse-smtp.{}', 587, 1),
	('constantcontact-mail.{}', 25, 0),
	('constantcontact-smtp.{}', 587, 1),
	('noreply-mail.{}', 25, 0),
	('noreply-smtp.{}', 587, 1),
	('no-reply-mail.{}', 25, 0),
	('no-reply-smtp.{}', 587, 1),
	('donotreply-mail.{}', 25, 0),
	('donotreply-smtp.{}', 587, 1),
	('do-not-reply-mail.{}', 25, 0),
	('do-not-reply-smtp.{}', 587, 1),
	('verify-mail.{}', 25, 0),
	('verification-mail.{}', 25, 0),
	('verification-smtp.{}', 587, 1),
	('confirm-mail.{}', 25, 0),
	('confirmation-mail.{}', 25, 0),
	('confirmation-smtp.{}', 587, 1),
	('activate-mail.{}', 25, 0),
	('activate-smtp.{}', 587, 1),
	('activation-mail.{}', 25, 0),
	('reset-mail.{}', 25, 0),
	('password-mail.{}', 25, 0),
	('recover-mail.{}', 25, 0),
	('recover-smtp.{}', 587, 1),
	('recovery-mail.{}', 25, 0),
	('welcome-mail.{}', 25, 0),
	('invite-mail.{}', 25, 0),
	('invite-smtp.{}', 587, 1),
	('notification-mail.{}', 25, 0),
	('notification-smtp.{}', 587, 1),
	('transaction-mail.{}', 25, 0),
	('transaction-smtp.{}', 587, 1),
	('mx-backup-mail.{}', 25, 0),
	('backup-mail-relay.{}', 25, 0),
	('backup-mx-relay.{}', 25, 0),
	('backup-smtp-relay.{}', 25, 0),
	('backup-relay-mail.{}', 25, 0),
	('backup-relay-mx.{}', 25, 0),
	('backup-relay-smtp.{}', 25, 0),
	('failover-mail-relay.{}', 25, 0),
	('failover-mx-relay.{}', 25, 0),
	('failover-smtp-relay.{}', 25, 0),
	('standby-mail-relay.{}', 25, 0),
	('standby-mx-relay.{}', 25, 0),
	('disaster-mx.{}', 25, 0),
	('disaster-recovery-mail.{}', 25, 0),
	('disaster-recovery-smtp.{}', 25, 0),
	('smtp-gb.{}', 25, 0),
	('smtp-us-east.{}', 25, 0),
	('smtp-us-west.{}', 25, 0),
	('smtp-us-central.{}', 25, 0),
	('smtp-eu-west.{}', 25, 0),
	('smtp-eu-east.{}', 25, 0),
	('smtp-eu-central.{}', 25, 0),
	('smtp-eu-north.{}', 25, 0),
	('smtp-eu-south.{}', 25, 0),
	('smtp-ap-south.{}', 25, 0),
	('smtp-ap-east.{}', 25, 0),
	('smtp-ap-southeast.{}', 25, 0),
	('smtp-ap-northeast.{}', 25, 0),
	('smtp-sa-east.{}', 25, 0),
	('smtp-sa-west.{}', 25, 0),
	('smtp-me-central.{}', 25, 0),
	('smtp-af-south.{}', 25, 0),
	('mail-us-east.{}', 25, 0),
	('mail-us-west.{}', 25, 0),
	('mail-eu-west.{}', 25, 0),
	('mail-eu-east.{}', 25, 0),
	('mail-ap-south.{}', 25, 0),
	('mail-ap-east.{}', 25, 0),
	('mx-us-east.{}', 25, 0),
	('mx-us-west.{}', 25, 0),
	('mx-eu-west.{}', 25, 0),
	('mx-eu-east.{}', 25, 0),
	('mx-ap-south.{}', 25, 0),
	('mx-ap-east.{}', 25, 0),
	('smtp-ipv4.{}', 25, 0),
	('smtp-ipv6.{}', 25, 0),
	('mail-ipv4.{}', 25, 0),
	('mail-ipv6.{}', 25, 0),
	('v4-mail.{}', 25, 0),
	('v6-mail.{}', 25, 0),
	('smtp-day.{}', 25, 0),
	('smtp-night.{}', 25, 0),
	('mail-day.{}', 25, 0),
	('mail-night.{}', 25, 0),
	('smtp-morning.{}', 25, 0),
	('smtp-evening.{}', 25, 0),
	('mail-morning.{}', 25, 0),
	('mail-evening.{}', 25, 0),
	('a-mail.{}', 25, 0),
	('b-mail.{}', 25, 0),
	('c-mail.{}', 25, 0),
	('d-mail.{}', 25, 0),
	('a-smtp.{}', 25, 0),
	('b-smtp.{}', 25, 0),
	('c-smtp.{}', 25, 0),
	('d-smtp.{}', 25, 0),
	('a-mx.{}', 25, 0),
	('b-mx.{}', 25, 0),
	('c-mx.{}', 25, 0),
	('d-mx.{}', 25, 0),
	('dovecot-smtp.{}', 25, 0),
	('cyrus-imap.{}', 143, 0),
	('mbox-mail.{}', 25, 0),
	('mdaemon.{}', 25, 0),
	('mdaemon-mail.{}', 25, 0),
	('mdaemon-smtp.{}', 25, 0),
	('vpopmail.{}', 25, 0),
	('vpopmail-mail.{}', 25, 0),
	('alias-mail.{}', 25, 0),
	('alias-smtp.{}', 25, 0),
	('virtual-mail-server.{}', 25, 0),
	('virtualmailserver.{}', 25, 0),
	('hostgator-mail.{}', 25, 0),
	('hostgator-smtp.{}', 25, 0),
	('bluehost-mail.{}', 25, 0),
	('bluehost-smtp.{}', 25, 0),
	('godaddy-mail.{}', 25, 0),
	('godaddy-smtp.{}', 25, 0),
	('namecheap-mail.{}', 25, 0),
	('namecheap-smtp.{}', 25, 0),
	('dreamhost-mail.{}', 25, 0),
	('dreamhost-smtp.{}', 25, 0),
	('siteground-mail.{}', 25, 0),
	('siteground-smtp.{}', 25, 0),
	('hostinger-mail.{}', 25, 0),
	('hostinger-smtp.{}', 25, 0),
	('ionos-mail.{}', 25, 0),
	('ionos-smtp.{}', 25, 0),
	('1and1-mail.{}', 25, 0),
	('1and1-smtp.{}', 25, 0),
	('ovh-mail.{}', 25, 0),
	('ovh-smtp.{}', 25, 0),
	('hetzner-mail.{}', 25, 0),
	('hetzner-smtp.{}', 25, 0),
	('digitalocean-mail.{}', 25, 0),
	('digitalocean-smtp.{}', 25, 0),
	('linode-mail.{}', 25, 0),
	('linode-smtp.{}', 25, 0),
	('vultr-mail.{}', 25, 0),
	('vultr-smtp.{}', 25, 0),
	('mx-enterprise.{}', 25, 0),
	('relay-enterprise.{}', 25, 0),
	('ent-relay.{}', 25, 0),
	('corp1-mail.{}', 25, 0),
	('corp2-mail.{}', 25, 0),
	('hq-relay.{}', 25, 0),
	('branch-mail.{}', 25, 0),
	('branch-smtp.{}', 25, 0),
	('branch-mx.{}', 25, 0),
	('central-mail.{}', 25, 0),
	('centralmail.{}', 25, 0),
	('central-smtp.{}', 25, 0),
	('centralsmtp.{}', 25, 0),
	('central-mx.{}', 25, 0),
	('centralmx.{}', 25, 0),
	('central-relay.{}', 25, 0),
	('centralrelay.{}', 25, 0),
	('main-mail-server.{}', 25, 0),
	('mainmailserver.{}', 25, 0),
	('primary-mail-server.{}', 25, 0),
	('primarymailserver.{}', 25, 0),
	('edge-relay.{}', 25, 0),
	('perimeter-mx.{}', 25, 0),
	('border-mx.{}', 25, 0),
	('front-relay.{}', 25, 0),
	('backend-mx.{}', 25, 0),
	('backend-relay.{}', 25, 0),
	('frontend-mx.{}', 25, 0),
	('frontend-relay.{}', 25, 0),
	('smarthost1.{}', 25, 0),
	('smarthost2.{}', 25, 0),
	('mx-relay-1.{}', 25, 0),
	('mx-relay-2.{}', 25, 0),
	('smtp-relay-1.{}', 25, 0),
	('smtp-relay-2.{}', 25, 0),
	('mail-relay-1.{}', 25, 0),
	('mail-relay-2.{}', 25, 0),
	('mxsrv.{}', 25, 0),
	('mxsrv1.{}', 25, 0),
	('mxsrv2.{}', 25, 0),
	('relaysrv.{}', 25, 0),
	('relaysrv1.{}', 25, 0),
	('relaysrv2.{}', 25, 0),
	('emailsrv.{}', 25, 0),
	('emailsrv1.{}', 25, 0),
	('emailsrv2.{}', 25, 0),
	('smtpserver1.{}', 25, 0),
	('smtpserver2.{}', 25, 0),
	('mailserver1.{}', 25, 0),
	('mailserver2.{}', 25, 0),
	('emailserver1.{}', 25, 0),
	('emailserver2.{}', 25, 0),
	('hostedmail.{}', 25, 0),
	('hostedsmtp.{}', 587, 1),
	('hosted-email.{}', 25, 0),
	('hostedemail.{}', 25, 0),
	('hosted-exchange.{}', 25, 0),
	('hostedexchange.{}', 25, 0),
	('managedmail.{}', 25, 0),
	('managedsmtp.{}', 587, 1),
	('managed-email.{}', 25, 0),
	('managedemail.{}', 25, 0),
	('dkim.{}', 25, 0),
	('spf.{}', 25, 0),
	('dmarc.{}', 25, 0),
	('authenticated.{}', 587, 1),
	('auth1-mail.{}', 25, 0),
	('auth2-mail.{}', 25, 0),
	('api-relay.{}', 25, 0),
	('apirelay.{}', 25, 0),
	('api-mx.{}', 25, 0),
	('apimx.{}', 25, 0),
	('svc-smtp.{}', 25, 0),
	('svcsmtp.{}', 25, 0),
	('smtpalt.{}', 2525, 0),
	('mailalt.{}', 2525, 0),
	('mx-alt.{}', 2525, 0),
	('mxalt.{}', 2525, 0),
	('relay-alt.{}', 2525, 0),
	('relayalt.{}', 2525, 0),
	('smtp-2.{}', 2525, 0),
	('mail-2.{}', 2525, 0),
	('altport.{}', 2525, 0),
	('custom-mail.{}', 2525, 0),
	('custom-smtp.{}', 2525, 0),
	('smtp-custom.{}', 2525, 0),
	('smtp-in1.{}', 587, 1),
	('smtp-in2.{}', 587, 1),
	('smtp-out1.{}', 587, 1),
	('smtp-out2.{}', 587, 1),
	('mx-in1.{}', 25, 0),
	('mx-in2.{}', 25, 0),
	('mx-out1.{}', 25, 0),
	('mx-out2.{}', 25, 0),
	('relay-in1.{}', 25, 0),
	('relay-in2.{}', 25, 0),
	('relay-out1.{}', 25, 0),
	('relay-out2.{}', 25, 0),
	('mx-relay-in.{}', 25, 0),
	('mx-relay-out.{}', 25, 0),
	('smtp-mx-in.{}', 25, 0),
	('smtp-mx-out.{}', 25, 0),
	('mail-mx-in.{}', 25, 0),
	('mail-mx-out.{}', 25, 0),
	('in-mail-relay.{}', 25, 0),
	('out-mail-relay.{}', 25, 0),
	('in-smtp-relay.{}', 25, 0),
	('out-smtp-relay.{}', 25, 0),
	('saas-mail.{}', 25, 0),
	('saas-smtp.{}', 587, 1),
	('cloud-mail.{}', 587, 1),
	('cloud-smtp.{}', 25, 0),
	('cloud-email.{}', 587, 1),
	('cloud-mx.{}', 25, 0),
	('cloud-relay.{}', 25, 0),
	('vps-mail.{}', 25, 0),
	('vps-smtp.{}', 587, 1),
	('dedicated-mail.{}', 25, 0),
	('dedicated-smtp.{}', 587, 1),
	('shared-mail.{}', 25, 0),
	('shared-smtp.{}', 587, 1),
	('reseller-mail.{}', 25, 0),
	('reseller-smtp.{}', 587, 1),
	('mail-edge.{}', 25, 0),
	('smtp-edge.{}', 25, 0),
	('mx-edge.{}', 25, 0),
	('relay-edge.{}', 25, 0),
	('edge-mail-1.{}', 25, 0),
	('edge-mail-2.{}', 25, 0),
	('edge-smtp-1.{}', 25, 0),
	('edge-smtp-2.{}', 25, 0),
	('edge-mx-1.{}', 25, 0),
	('edge-mx-2.{}', 25, 0),
	('cdn-mail.{}', 25, 0),
	('cdn-smtp.{}', 25, 0),
	('cdn-mx.{}', 25, 0),
	('cdn-relay.{}', 25, 0),
	('pop-mail.{}', 25, 0),
	('pop-smtp.{}', 25, 0),
	('proxy-mail.{}', 25, 0),
	('proxy-smtp.{}', 25, 0),
	('proxy-mx.{}', 25, 0),
	('proxy-relay.{}', 25, 0),
	('forward-mail-1.{}', 25, 0),
	('forward-mail-2.{}', 25, 0),
	('forwarder-smtp.{}', 25, 0),
	('redirect-smtp.{}', 25, 0),
	('mail-smtp-relay-gw.{}', 25, 0),
	('smtp-mail-relay-gw.{}', 25, 0),
	('relay-mail-smtp-gw.{}', 25, 0),
	('gw-mail-smtp-relay.{}', 25, 0),
	('gw-smtp-mail-relay.{}', 25, 0),
	('relay-gw-mail-smtp.{}', 25, 0),
	('mail-out-relay-smtp.{}', 25, 0),
	('smtp-out-relay-mail.{}', 25, 0),
	('relay-out-mail-smtp.{}', 25, 0),
	('smtp-in-relay-mail.{}', 25, 0),
	('mail-in-relay-smtp.{}', 25, 0),
	('relay-in-mail-smtp.{}', 25, 0),
	('primary-mail-relay.{}', 25, 0),
	('primary-smtp-relay.{}', 25, 0),
	('secondary-mail-relay.{}', 25, 0),
	('secondary-smtp-relay.{}', 25, 0),
	('backup-mail-relay-gw.{}', 25, 0),
	('backup-smtp-relay-gw.{}', 25, 0),
	('smtp-tls.{}', 25, 0),
	('mail-tls.{}', 25, 0),
	('tls-relay.{}', 587, 1),
	('tls-relay.{}', 25, 0),
	('ssl-relay.{}', 465, 1),
	('ssl-relay.{}', 25, 0),
	('tls-mx.{}', 587, 1),
	('ssl-mx.{}', 465, 1),
	('smtp3a.{}', 25, 0),
	('smtp3b.{}', 25, 0),
	('mail3a.{}', 25, 0),
	('mail3b.{}', 25, 0),
	('relay1a.{}', 25, 0),
	('relay1b.{}', 25, 0),
	('relay2a.{}', 25, 0),
	('relay2b.{}', 25, 0),
	('mx3a.{}', 25, 0),
	('mx3b.{}', 25, 0),
	('esmtp1.{}', 25, 0),
	('esmtp2.{}', 25, 0),
	('esmtp-relay.{}', 587, 1),
	('smtps1.{}', 465, 1),
	('smtps2.{}', 465, 1),
	('smtps1.{}', 587, 1),
	('smtps2.{}', 587, 1),
	('simple-mail.{}', 25, 0),
	('simple-mail.{}', 587, 1),
	('simplemail.{}', 25, 0),
	('simplemail.{}', 587, 1),
	('simple-smtp.{}', 25, 0),
	('simple-smtp.{}', 587, 1),
	('simplesmtp.{}', 25, 0),
	('simplesmtp.{}', 587, 1),
	('simple-mx.{}', 25, 0),
	('simplemx.{}', 25, 0),
	('m-smtp.{}', 25, 0),
	('m-mail.{}', 25, 0),
	('e-smtp.{}', 25, 0),
	('e-mail-server.{}', 25, 0),
	('mail-s.{}', 25, 0),
	('mail-m.{}', 25, 0),
	('mail-e.{}', 25, 0),
	('smtp-s.{}', 25, 0),
	('smtp-m.{}', 25, 0),
	('smtp-e.{}', 25, 0),
	('mx-s.{}', 25, 0),
	('mx-m.{}', 25, 0),
	('mx-e.{}', 25, 0),
	('smtp-mail-out.{}', 587, 1),
	('smtp-mail-out.{}', 25, 0),
	('smtp-mail-in.{}', 25, 0),
	('mail-smtp-out.{}', 587, 1),
	('mail-smtp-out.{}', 25, 0),
	('mail-smtp-in.{}', 25, 0),
	('out-mail-smtp-relay.{}', 587, 1),
	('out-mail-smtp-relay.{}', 25, 0),
	('in-mail-smtp-relay.{}', 25, 0),
	('mail-relay-out.{}', 587, 1),
	('smtp-relay-out.{}', 587, 1),
	('relay-mail-out.{}', 587, 1),
	('relay-mail-in.{}', 25, 0),
	('mail-hosted-1.{}', 25, 0),
	('mail-hosted-2.{}', 25, 0),
	('smtp-hosted-1.{}', 587, 1),
	('smtp-hosted-2.{}', 587, 1),
	('secure-mail-relay.{}', 465, 1),
	('secure-smtp-relay.{}', 465, 1),
	('secure-mx-relay.{}', 465, 1),
	('secure-email-relay.{}', 465, 1),
	('secure-relay-mail.{}', 465, 1),
	('secure-relay-smtp.{}', 465, 1),
	('secure-relay-mx.{}', 465, 1),
	('secure-relay-email.{}', 465, 1),
	('secure1-mail.{}', 465, 1),
	('secure2-mail.{}', 465, 1),
	('secure1-smtp.{}', 465, 1),
	('secure2-smtp.{}', 465, 1),
	('secure1-mx.{}', 465, 1),
	('secure2-mx.{}', 465, 1),
	('smtp-mail-mx-relay.{}', 25, 0),
	('mail-smtp-mx-relay.{}', 25, 0),
	('mx-mail-smtp-relay.{}', 25, 0),
	('relay-mx-mail-smtp.{}', 25, 0),
	('smtp-relay-mail-mx.{}', 25, 0),
	('mail-relay-smtp-mx.{}', 25, 0),
	('mx-relay-smtp-mail.{}', 25, 0),
	('relay-mx-smtp-mail.{}', 25, 0),
	('smtp-br.{}', 25, 0),
	('smtp-es.{}', 25, 0),
	('smtp-it.{}', 25, 0),
	('smtp-ru.{}', 25, 0),
	('smtp-nl.{}', 25, 0),
	('smtp-se.{}', 25, 0),
	('smtp-no.{}', 25, 0),
	('smtp-fi.{}', 25, 0),
	('smtp-dk.{}', 25, 0),
	('smtp-pl.{}', 25, 0),
	('smtp-at.{}', 25, 0),
	('smtp-ch.{}', 25, 0),
	('smtp-be.{}', 25, 0),
	('smtp-pt.{}', 25, 0),
	('smtp-gr.{}', 25, 0),
	('smtp-tr.{}', 25, 0),
	('smtp-kr.{}', 25, 0),
	('smtp-tw.{}', 25, 0),
	('smtp-hk.{}', 25, 0),
	('smtp-sg.{}', 25, 0),
	('smtp-my.{}', 25, 0),
	('smtp-th.{}', 25, 0),
	('smtp-vn.{}', 25, 0),
	('smtp-id.{}', 25, 0),
	('smtp-ph.{}', 25, 0),
	('smtp-nz.{}', 25, 0),
	('smtp-za.{}', 25, 0),
	('smtp-ae.{}', 25, 0),
	('smtp-sa.{}', 25, 0),
	('smtp-il.{}', 25, 0),
	('smtp-ar.{}', 25, 0),
	('smtp-cl.{}', 25, 0),
	('smtp-co.{}', 25, 0),
	('smtp-pe.{}', 25, 0),
	('mx-gateway-1.{}', 25, 0),
	('mx-gateway-2.{}', 25, 0),
	('relay-gateway-1.{}', 25, 0),
	('relay-gateway-2.{}', 25, 0),
	('email-gateway-1.{}', 25, 0),
	('email-gateway-2.{}', 25, 0),
	('emailgw1.{}', 25, 0),
	('emailgw2.{}', 25, 0),
	('smtp465.{}', 465, 1),
	('smtp587.{}', 587, 1),
	('mail25.{}', 25, 0),
	('mail465.{}', 465, 1),
	('mail587.{}', 587, 1),
	('mx25.{}', 25, 0),
	('mx465.{}', 465, 1),
	('mx587.{}', 587, 1),
	('fast-mail.{}', 25, 0),
	('fastmail.{}', 25, 0),
	('fast-smtp.{}', 587, 1),
	('fastsmtp.{}', 587, 1),
	('quick-mail.{}', 25, 0),
	('quickmail.{}', 25, 0),
	('quick-smtp.{}', 587, 1),
	('quicksmtp.{}', 587, 1),
	('rapid-mail.{}', 25, 0),
	('rapidmail.{}', 25, 0),
	('rapid-smtp.{}', 587, 1),
	('rapidsmtp.{}', 587, 1),
	('speed-mail.{}', 25, 0),
	('speedmail.{}', 25, 0),
	('speed-smtp.{}', 587, 1),
	('speedsmtp.{}', 587, 1),
	('turbo-mail.{}', 25, 0),
	('turbomail.{}', 25, 0),
	('turbo-smtp.{}', 587, 1),
	('turbosmtp.{}', 587, 1),
	('relayserver.{}', 25, 0),
	('mxserver.{}', 25, 0),
	('smtp03.{}', 25, 0),
	('smtp04.{}', 25, 0),
	('smtp.secureserver.{}', 25, 0),
	('smtpout.{}', 465, 1),
	('smtp2go.{}', 25, 0),
	('smtpin.{}', 465, 1),
	('smtps.{}', 587, 1),
	('smtpclient.{}', 25, 0),
	('smtpsend.{}', 25, 0),
	('smtp.mailserver.{}', 25, 0),
	('smtp.msg.{}', 25, 0),
	('smtp.mess.{}', 25, 0),
	('smtp.test.{}', 25, 0),
	('smtp.dev.{}', 25, 0),
	('smtp.staging.{}', 25, 0),
	('smtp.prod.{}', 25, 0),
	('smtp.office.{}', 587, 1),
	('smtp.biz.{}', 587, 1),
	('smtp.cloud.{}', 587, 1),
	('smtp.web.{}', 25, 0),
	('smtp.host.{}', 25, 0),
	('smtp.server.{}', 25, 0),
	('smtp.net.{}', 25, 0),
	('smtp.link.{}', 25, 0),
	('smtp.systems.{}', 25, 0),
	('smtp.services.{}', 25, 0),
	('smtp.solutions.{}', 25, 0),
	('smtp.tech.{}', 587, 1),
	('smtp.io.{}', 587, 1),
	('mail03.{}', 25, 0),
	('mail04.{}', 25, 0),
	('mail05.{}', 25, 0),
	('mail-in.{}', 465, 1),
	('mailout.{}', 587, 1),
	('mail-gate.{}', 25, 0),
	('mailgate.{}', 25, 0),
	('mailscan.{}', 25, 0),
	('mail-scanner.{}', 25, 0),
	('mailserv.{}', 25, 0),
	('mail-serv.{}', 25, 0),
	('mailsend.{}', 25, 0),
	('mailrecv.{}', 25, 0),
	('mail-recv.{}', 25, 0),
	('mailsecure.{}', 465, 1),
	('mail-secure.{}', 465, 1),
	('mailauth.{}', 587, 1),
	('mail-auth.{}', 587, 1),
	('mail.msg.{}', 25, 0),
	('mail.web.{}', 25, 0),
	('mail.cloud.{}', 587, 1),
	('mail.office.{}', 587, 1),
	('mail.biz.{}', 587, 1),
	('mail.host.{}', 25, 0),
	('mail.net.{}', 25, 0),
	('mail.io.{}', 587, 1),
	('mail.link.{}', 25, 0),
	('mx-in.{}', 25, 0),
	('mxin.{}', 25, 0),
	('mx-out.{}', 25, 0),
	('mxout.{}', 25, 0),
	('mx-gateway.{}', 25, 0),
	('mxgateway.{}', 25, 0),
	('mx-proxy.{}', 25, 0),
	('mxproxy.{}', 25, 0),
	('mx-filter.{}', 25, 0),
	('mxfilter.{}', 25, 0),
	('mx-server.{}', 25, 0),
	('mx-host.{}', 25, 0),
	('mxhost.{}', 25, 0),
	('mx-secure.{}', 465, 1),
	('mxsecure.{}', 465, 1),
	('mx-ssl.{}', 465, 1),
	('mxssl.{}', 465, 1),
	('mxauth.{}', 587, 1),
	('mx-auth.{}', 587, 1),
	('mx-in.{}', 465, 1),
	('mx-out.{}', 587, 1),
	('mx.cloud.{}', 587, 1),
	('mx.corp.{}', 25, 0),
	('mx.internal.{}', 25, 0),
	('mx.io.{}', 587, 1),
	('relay03.{}', 25, 0),
	('relay-mx.{}', 25, 0),
	('relaymx.{}', 25, 0),
	('relay-gw.{}', 25, 0),
	('relaygw.{}', 25, 0),
	('relay-secure.{}', 465, 1),
	('relaysecure.{}', 465, 1),
	('relay-ssl.{}', 465, 1),
	('relayssl.{}', 465, 1),
	('relay-auth.{}', 587, 1),
	('relayauth.{}', 587, 1),
	('relay.internal.{}', 25, 0),
	('relay.corp.{}', 25, 0),
	('relay.dmz.{}', 25, 0),
	('relay.email.{}', 25, 0),
	('relay.host.{}', 25, 0),
	('relay.net.{}', 25, 0),
	('exchange1.{}', 25, 0),
	('exchange2.{}', 25, 0),
	('exchange01.{}', 25, 0),
	('exchange02.{}', 25, 0),
	('exchangemail.{}', 25, 0),
	('exchangesmtp.{}', 25, 0),
	('exchangerelay.{}', 25, 0),
	('exchange-mx.{}', 25, 0),
	('exchangemx.{}', 25, 0),
	('exchange-secure.{}', 465, 1),
	('exchange-ssl.{}', 465, 1),
	('exchange-auth.{}', 587, 1),
	('exchange.internal.{}', 25, 0),
	('exchange.corp.{}', 25, 0),
	('exchange.mail.{}', 25, 0),
	('exch01.{}', 25, 0),
	('exch02.{}', 25, 0),
	('post1.{}', 25, 0),
	('post2.{}', 25, 0),
	('post-relay.{}', 25, 0),
	('postrelay.{}', 25, 0),
	('post-box.{}', 25, 0),
	('postfix1.{}', 25, 0),
	('postfix2.{}', 25, 0),
	('gw1.{}', 25, 0),
	('gw2.{}', 25, 0),
	('gw01.{}', 25, 0),
	('gw02.{}', 25, 0),
	('gw-mail.{}', 25, 0),
	('gwmail.{}', 25, 0),
	('gw-smtp.{}', 25, 0),
	('gwsmtp.{}', 25, 0),
	('gw-relay.{}', 25, 0),
	('gwrelay.{}', 25, 0),
	('send1.{}', 25, 0),
	('send2.{}', 25, 0),
	('send01.{}', 25, 0),
	('send02.{}', 25, 0),
	('sendout.{}', 25, 0),
	('send-out.{}', 25, 0),
	('sender.{}', 25, 0),
	('sending.{}', 25, 0),
	('smtp-send.{}', 587, 1),
	('mail-send.{}', 587, 1),
	('out1.{}', 25, 0),
	('out2.{}', 25, 0),
	('outbound1.{}', 25, 0),
	('outbound2.{}', 25, 0),
	('outmail.{}', 25, 0),
	('outsmtp.{}', 25, 0),
	('out-relay.{}', 25, 0),
	('outrelay.{}', 25, 0),
	('outboundmail.{}', 25, 0),
	('outbound-relay.{}', 25, 0),
	('outboundrelay.{}', 25, 0),
	('outbound-smtp.{}', 587, 1),
	('outbound-smtp.{}', 25, 0),
	('in.{}', 25, 0),
	('in1.{}', 25, 0),
	('in2.{}', 25, 0),
	('inbound1.{}', 25, 0),
	('inbound2.{}', 25, 0),
	('inmail.{}', 25, 0),
	('insmtp.{}', 25, 0),
	('in-relay.{}', 25, 0),
	('inrelay.{}', 25, 0),
	('inboundmail.{}', 25, 0),
	('inbound-relay.{}', 25, 0),
	('inboundrelay.{}', 25, 0),
	('pops.{}', 995, 1),
	('pop3s.{}', 995, 1),
	('pop-mail.{}', 110, 0),
	('popmail.{}', 110, 0),
	('imaps.{}', 993, 1),
	('imap-mail.{}', 143, 0),
	('imapmail.{}', 143, 0),
	('imap4.{}', 143, 0),
	('imap4.{}', 993, 1),
	('webmail.{}', 25, 0),
	('owa.{}', 25, 0),
	('googlemail.{}', 25, 0),
	('gmail.{}', 25, 0),
	('aspmx.{}', 25, 0),
	('aspmx1.{}', 25, 0),
	('aspmx2.{}', 25, 0),
	('aspmx3.{}', 25, 0),
	('outlook.{}', 25, 0),
	('outlook.{}', 587, 1),
	('office365.{}', 587, 1),
	('microsoft.{}', 25, 0),
	('zoho.{}', 25, 0),
	('yahoo.{}', 25, 0),
	('hotmail.{}', 25, 0),
	('icloud.{}', 25, 0),
	('proton.{}', 25, 0),
	('protonmail.{}', 25, 0),
	('yandex.{}', 25, 0),
	('mandrill.{}', 25, 0),
	('mailchimp.{}', 25, 0),
	('amazonses.{}', 25, 0),
	('postmark.{}', 25, 0),
	('sparkpost.{}', 25, 0),
	('barracuda.{}', 587, 1),
	('mimecast.{}', 587, 1),
	('symantec.{}', 25, 0),
	('trendmicro.{}', 25, 0),
	('mcafee.{}', 25, 0),
	('antivirus.{}', 25, 0),
	('anti-virus.{}', 25, 0),
	('virusfilter.{}', 25, 0),
	('virus-filter.{}', 25, 0),
	('messagefilter.{}', 25, 0),
	('message-filter.{}', 25, 0),
	('mta3.{}', 25, 0),
	('mta01.{}', 25, 0),
	('mta02.{}', 25, 0),
	('mta03.{}', 25, 0),
	('mta-1.{}', 25, 0),
	('mta-2.{}', 25, 0),
	('mta-3.{}', 25, 0),
	('msg.{}', 25, 0),
	('messaging.{}', 25, 0),
	('message.{}', 25, 0),
	('messages.{}', 25, 0),
	('msg-server.{}', 25, 0),
	('msgserver.{}', 25, 0),
	('msg-gw.{}', 25, 0),
	('msggw.{}', 25, 0),
	('message-gateway.{}', 25, 0),
	('messagegateway.{}', 25, 0),
	('messaging-server.{}', 25, 0),
	('messagingserver.{}', 25, 0),
	('smtp-cn.{}', 25, 0),
	('smtp-north.{}', 25, 0),
	('smtp-south.{}', 25, 0),
	('smtp-central.{}', 25, 0),
	('smtp05.{}', 25, 0),
	('smtp06.{}', 25, 0),
	('smtp07.{}', 25, 0),
	('smtp08.{}', 25, 0),
	('smtp09.{}', 25, 0),
	('smtp10.{}', 587, 1),
	('mail06.{}', 25, 0),
	('mail07.{}', 25, 0),
	('mail08.{}', 25, 0),
	('mail09.{}', 25, 0),
	('mail10.{}', 587, 1),
	('mx10.{}', 587, 1),
	('email1.{}', 25, 0),
	('email2.{}', 25, 0),
	('email01.{}', 25, 0),
	('email02.{}', 25, 0),
	('e-mail.{}', 25, 0),
	('e-mail.{}', 587, 1),
	('emailer.{}', 25, 0),
	('emailer1.{}', 25, 0),
	('emailer2.{}', 25, 0),
	('email-relay.{}', 25, 0),
	('emailrelay.{}', 25, 0),
	('emailsmtp.{}', 587, 1),
	('email-gw.{}', 25, 0),
	('emailgw.{}', 25, 0),
	('email-mx.{}', 25, 0),
	('emailmx.{}', 25, 0),
	('email-secure.{}', 465, 1),
	('emailsecure.{}', 465, 1),
	('email-ssl.{}', 465, 1),
	('emailssl.{}', 465, 1),
	('email-auth.{}', 587, 1),
	('emailauth.{}', 587, 1),
	('email-filter.{}', 25, 0),
	('emailfilter.{}', 25, 0),
	('email-scan.{}', 25, 0),
	('emailscan.{}', 25, 0),
	('email-cloud.{}', 587, 1),
	('emailcloud.{}', 587, 1),
	('email-mail.{}', 25, 0),
	('postal.{}', 25, 0),
	('postman.{}', 25, 0),
	('courier.{}', 25, 0),
	('bulk.{}', 25, 0),
	('news.{}', 25, 0),
	('nntp.{}', 119, 0),
	('newsgroup.{}', 119, 0),
	('usenet.{}', 119, 0),
	('list.{}', 25, 0),
	('lists.{}', 25, 0),
	('mailinglist.{}', 25, 0),
	('mailing-list.{}', 25, 0),
	('mailman.{}', 25, 0),
	('mailqueue1.{}', 25, 0),
	('mailqueue2.{}', 25, 0),
	('queue.{}', 25, 0),
	('mbox.{}', 25, 0),
	('maildir.{}', 25, 0),
	('box.{}', 25, 0),
	('in-box.{}', 25, 0),
	('out-box.{}', 25, 0),
	('forward.{}', 25, 0),
	('forwarder.{}', 25, 0),
	('redirect.{}', 25, 0),
	('bounce.{}', 25, 0),
	('bouncer.{}', 25, 0),
	('internal-relay.{}', 25, 0),
	('internalrelay.{}', 25, 0),
	('corp-relay.{}', 25, 0),
	('corprelay.{}', 25, 0),
	('corp-mx.{}', 25, 0),
	('corpmx.{}', 25, 0),
	('local-mail.{}', 25, 0),
	('localmail.{}', 25, 0),
	('local-smtp.{}', 25, 0),
	('localsmtp.{}', 25, 0),
	('local-relay.{}', 25, 0),
	('localrelay.{}', 25, 0),
	('dmzmail.{}', 25, 0),
	('dmzsmtp.{}', 25, 0),
	('dmz-relay.{}', 25, 0),
	('dmzrelay.{}', 25, 0),
	('ns.{}', 25, 0),
	('ns1.{}', 25, 0),
	('ns2.{}', 25, 0),
	('www-mail.{}', 25, 0),
	('wwwmail.{}', 25, 0),
	('web-mail.{}', 25, 0),
	('web-smtp.{}', 587, 1),
	('webrtc-mail.{}', 25, 0),
	('apis.{}', 25, 0),
	('virtualmin.{}', 25, 0),
	('webmin.{}', 25, 0),
	('postfixadmin.{}', 25, 0),
	('exim-mail.{}', 25, 0),
	('msexchange.{}', 25, 0),
	('ms-exchange.{}', 25, 0),
	('lotus.{}', 25, 0),
	('lotusnotes.{}', 25, 0),
	('lotus-notes.{}', 25, 0),
	('domino.{}', 25, 0),
	('notes.{}', 25, 0),
	('groupwise.{}', 25, 0),
	('kerio.{}', 25, 0),
	('exim4.{}', 587, 1),
	('smtp.out.{}', 587, 1),
	('mail.out.{}', 25, 0),
	('out.mail.{}', 25, 0),
	('smtp.relay.{}', 587, 1),
	('smtp.mx.{}', 587, 1),
	('smtp.gw.{}', 25, 0),
	('mail.gw.{}', 25, 0),
	('gw.mail.{}', 25, 0),
	('smtp.gw.{}', 587, 1),
	('smtp1.smtp.{}', 25, 0),
	('smtp2.smtp.{}', 25, 0),
	('relay1.relay.{}', 25, 0),
	('relay2.relay.{}', 25, 0),
	('mx1.mx.{}', 25, 0),
	('mx2.mx.{}', 25, 0),
	('secure.{}', 465, 1),
	('secure.{}', 587, 1),
	('secure1.{}', 465, 1),
	('secure2.{}', 465, 1),
	('secure-mail.{}', 587, 1),
	('encrypt-mail.{}', 465, 1),
	('encryptmail.{}', 465, 1),
	('encrypted.{}', 465, 1),
	('crypto-mail.{}', 465, 1),
	('start-tls.{}', 587, 1),
	('starttls-mail.{}', 587, 1),
	('smtp-mail-relay.{}', 25, 0),
	('mail-smtp-relay.{}', 25, 0),
	('relay-mail-smtp.{}', 25, 0),
	('smtp-relay-mail.{}', 25, 0),
	('mail-relay-smtp.{}', 25, 0),
	('smtp-mx-relay.{}', 25, 0),
	('mx-smtp-relay.{}', 25, 0),
	('smtp-out-mail.{}', 587, 1),
	('out-mail-smtp.{}', 587, 1),
	('mail-out-smtp.{}', 587, 1),
	('smtp-in-mail.{}', 25, 0),
	('in-mail-smtp.{}', 25, 0),
	('mail-in-smtp.{}', 25, 0),
	('email-smtp-relay.{}', 25, 0),
	('smtp-email-relay.{}', 25, 0),
	('relay-email-smtp.{}', 25, 0),
	('smtp-1.{}', 25, 0),
	('smtp-2.{}', 25, 0),
	('smtp-3.{}', 25, 0),
	('smtp-4.{}', 25, 0),
	('smtp-5.{}', 25, 0),
	('smtp-01.{}', 25, 0),
	('smtp-02.{}', 25, 0),
	('smtp-03.{}', 25, 0),
	('smtp-04.{}', 25, 0),
	('smtp-05.{}', 25, 0),
	('mail-1.{}', 25, 0),
	('mail-2.{}', 25, 0),
	('mail-3.{}', 25, 0),
	('mail-4.{}', 25, 0),
	('mail-5.{}', 25, 0),
	('mail-01.{}', 25, 0),
	('mail-02.{}', 25, 0),
	('mail-03.{}', 25, 0),
	('mail-04.{}', 25, 0),
	('mail-05.{}', 25, 0),
	('mx-1.{}', 25, 0),
	('mx-2.{}', 25, 0),
	('mx-3.{}', 25, 0),
	('mx-4.{}', 25, 0),
	('mx-5.{}', 25, 0),
	('mx-01.{}', 25, 0),
	('mx-02.{}', 25, 0),
	('mx-03.{}', 25, 0),
	('mx-04.{}', 25, 0),
	('mx-05.{}', 25, 0),
	('relay-1.{}', 25, 0),
	('relay-2.{}', 25, 0),
	('relay-3.{}', 25, 0),
	('relay-01.{}', 25, 0),
	('relay-02.{}', 25, 0),
	('relay-03.{}', 25, 0),
	('postmail.{}', 587, 1),
	('mailpost.{}', 25, 0),
	('mail-post.{}', 25, 0),
	('mail-daemon.{}', 25, 0),
	('mailerdaemon.{}', 25, 0),
	('postmaster-mail.{}', 25, 0),
	('webmaster-mail.{}', 25, 0),
	('admin-mail.{}', 25, 0),
	('adminmail.{}', 25, 0),
	('admin-smtp.{}', 587, 1),
	('adminsmtp.{}', 587, 1),
	('m.{}', 25, 0),
	('m1.{}', 25, 0),
	('m2.{}', 25, 0),
	('m01.{}', 25, 0),
	('m02.{}', 25, 0),
	('e.{}', 25, 0),
	('e1.{}', 25, 0),
	('e2.{}', 25, 0),
	('e01.{}', 25, 0),
	('e02.{}', 25, 0),
	('s.{}', 25, 0),
	('s1.{}', 25, 0),
	('s2.{}', 25, 0),
	('s01.{}', 25, 0),
	('s02.{}', 25, 0),
	('mx0.{}', 25, 0),
	('mail0.{}', 25, 0),
	('smtp0.{}', 25, 0),
	('relay0.{}', 25, 0),
	('email0.{}', 25, 0),
	('smtp-d.{}', 25, 0),
	('mail-d.{}', 25, 0),
	('mx-d.{}', 25, 0),
	('relay-c.{}', 25, 0),
	('relay-d.{}', 25, 0),
	('messaging.{}', 587, 1),
	('alert-mail.{}', 587, 1),
	('alertmail.{}', 587, 1),
	('automation-mail.{}', 587, 1),
	('batch-mail.{}', 25, 0),
	('batchmail.{}', 25, 0),
	('erp-mail.{}', 25, 0),
	('erpmail.{}', 25, 0),
	('accounting-mail.{}', 25, 0),
	('accountingmail.{}', 25, 0),
	('backupmail.{}', 25, 0),
	('backupsmtp.{}', 25, 0),
	('backupmx.{}', 25, 0),
	('fallbackmail.{}', 25, 0),
	('fallbacksmtp.{}', 25, 0),
	('fallbackmx.{}', 25, 0),
	('failovermail.{}', 25, 0),
	('failoversmtp.{}', 25, 0),
	('failovermx.{}', 25, 0),
	('standbymail.{}', 25, 0),
	('standbysmtp.{}', 25, 0),
	('standbymx.{}', 25, 0),
	('mirrormail.{}', 25, 0),
	('mirrorsmtp.{}', 25, 0),
	('mirrormx.{}', 25, 0),
	('replicamail.{}', 25, 0),
	('replicasmtp.{}', 25, 0),
	('replicamx.{}', 25, 0),
	('devmail.{}', 25, 0),
	('devsmtp.{}', 25, 0),
	('dev-mx.{}', 25, 0),
	('devmx.{}', 25, 0),
	('staging-mail.{}', 25, 0),
	('stagingmail.{}', 25, 0),
	('staging-smtp.{}', 25, 0),
	('stagingsmtp.{}', 25, 0),
	('staging-mx.{}', 25, 0),
	('stagingmx.{}', 25, 0),
	('prodmail.{}', 25, 0),
	('prodsmtp.{}', 587, 1),
	('prod-mx.{}', 25, 0),
	('prodmx.{}', 25, 0),
	('qa-mail.{}', 25, 0),
	('qamail.{}', 25, 0),
	('qa-smtp.{}', 25, 0),
	('qasmtp.{}', 25, 0),
	('testmail.{}', 25, 0),
	('testsmtp.{}', 25, 0),
	('test-mx.{}', 25, 0),
	('testmx.{}', 25, 0),
	('demo-mail.{}', 25, 0),
	('demomail.{}', 25, 0),
	('demo-smtp.{}', 25, 0),
	('demosmtp.{}', 25, 0),
	('sandbox-mail.{}', 25, 0),
	('sandboxmail.{}', 25, 0),
	('sandbox-smtp.{}', 25, 0),
	('sandboxsmtp.{}', 25, 0),
	('sandbox-mx.{}', 25, 0),
	('sandboxmx.{}', 25, 0),
	('smtp-alpha.{}', 25, 0),
	('smtp-beta.{}', 25, 0),
	('smtp-gamma.{}', 25, 0),
	('mail-alpha.{}', 25, 0),
	('mail-beta.{}', 25, 0),
	('mx-alpha.{}', 25, 0),
	('mx-beta.{}', 25, 0),
	('relay-alpha.{}', 25, 0),
	('relay-beta.{}', 25, 0),
	('smtp-oci.{}', 25, 0),
	('smtp-do.{}', 25, 0),
	('mx-aws.{}', 25, 0),
	('mx-azure.{}', 25, 0),
	('mx-gcp.{}', 25, 0),
	('relay-aws.{}', 25, 0),
	('relay-azure.{}', 25, 0),
	('relay-gcp.{}', 25, 0),
	('e1.mail.{}', 25, 0),
	('e2.mail.{}', 25, 0),
	('m1.mail.{}', 25, 0),
	('m2.mail.{}', 25, 0),
	('n1.mail.{}', 25, 0),
	('n2.mail.{}', 25, 0),
	('p1.mail.{}', 25, 0),
	('p2.mail.{}', 25, 0),
	('r1.mail.{}', 25, 0),
	('r2.mail.{}', 25, 0),
	('postino.{}', 25, 0),
	('postino1.{}', 25, 0),
	('buzon.{}', 25, 0),
	('casella.{}', 25, 0),
	('boite.{}', 25, 0),
	('inbox-server.{}', 25, 0),
	('inboxserver.{}', 25, 0),
	('outbox-server.{}', 25, 0),
	('outboxserver.{}', 25, 0),
	('mail-proxy1.{}', 25, 0),
	('mail-proxy2.{}', 25, 0),
	('smtp-proxy1.{}', 25, 0),
	('smtp-proxy2.{}', 25, 0),
	('email-proxy1.{}', 25, 0),
	('email-proxy2.{}', 25, 0),
	('mx3.mail.{}', 25, 0),
	('mx4.mail.{}', 25, 0),
	('smtp3.mail.{}', 25, 0),
	('smtp4.mail.{}', 25, 0),
	('relay3.mail.{}', 25, 0),
	('relay4.mail.{}', 25, 0),
	('mail.mail1.{}', 25, 0),
	('mail.mail2.{}', 25, 0),
	('smtp.smtp1.{}', 25, 0),
	('smtp.smtp2.{}', 25, 0),
	('mx.mx1.{}', 25, 0),
	('mx.mx2.{}', 25, 0),
	('relay.relay1.{}', 25, 0),
	('relay.relay2.{}', 25, 0),
	('securemail.{}', 587, 1),
	('secure-email.{}', 587, 1),
	('secureemail.{}', 587, 1),
	('ssl-email.{}', 465, 1),
	('sslemail.{}', 465, 1),
	('tls-email.{}', 587, 1),
	('tlsemail.{}', 587, 1),
	('emailssl.{}', 587, 1),
	('emailtls.{}', 587, 1),
	('mail-ssl.{}', 587, 1),
	('mailssl.{}', 587, 1),
	('smtp-ssl.{}', 587, 1),
	('smtpssl.{}', 587, 1),
	('smtpsecure.{}', 587, 1),
	('nodemail.{}', 25, 0),
	('vipmail.{}', 25, 0),
	('virtualsmtp.{}', 587, 0),
]

IMAP_SERVERS = [
    ("imap.{}", 993, 1),
    ("imap.{}", 143, 0),
    ("imap4.{}", 993, 1),
    ("imap4.{}", 143, 0),
    ("imap1.{}", 993, 1),
    ("imap1.{}", 143, 0),
    ("imap2.{}", 993, 1),
    ("imap2.{}", 143, 0),
    ("imap01.{}", 993, 1),
    ("imap01.{}", 143, 0),
    ("imap02.{}", 993, 1),
    ("imap02.{}", 143, 0),
    ("imap-mail.{}", 993, 1),
    ("imap-mail.{}", 143, 0),
    ("imapmail.{}", 993, 1),
    ("imapmail.{}", 143, 0),
    ("secure-imap.{}", 993, 1),
    ("secure-imap.{}", 143, 0),
    ("mail.{}", 993, 1),
    ("mail.{}", 143, 0),
    ("mail1.{}", 993, 1),
    ("mail1.{}", 143, 0),
    ("mail2.{}", 993, 1),
    ("mail2.{}", 143, 0),
    ("mail01.{}", 993, 1),
    ("mail01.{}", 143, 0),
    ("mail02.{}", 993, 1),
    ("mail02.{}", 143, 0),
    ("mailhost.{}", 993, 1),
    ("mailhost.{}", 143, 0),
    ("mailserver.{}", 993, 1),
    ("mailserver.{}", 143, 0),
    ("email.{}", 993, 1),
    ("email.{}", 143, 0),
    ("inbox.{}", 993, 1),
    ("inbox.{}", 143, 0),
    ("mailbox.{}", 993, 1),
    ("mailbox.{}", 143, 0),
    ("webmail.{}", 993, 1),
    ("webmail.{}", 143, 0),
    ("securemail.{}", 993, 1),
    ("securemail.{}", 143, 0),
    ("mx.{}", 993, 1),
    ("mx.{}", 143, 0),
    ("mx1.{}", 993, 1),
    ("mx1.{}", 143, 0),
    ("mx2.{}", 993, 1),
    ("mx2.{}", 143, 0),
    ("exchange.{}", 993, 1),
    ("exchange.{}", 143, 0),
    ("owa.{}", 993, 1),
    ("owa.{}", 143, 0),
]

POP3_SERVERS = [
    ("pop.{}", 995, 1),
    ("pop.{}", 110, 0),
    ("pop3.{}", 995, 1),
    ("pop3.{}", 110, 0),
    ("mail.{}", 995, 1),
    ("mail.{}", 110, 0),
    ("pop3s.{}", 995, 1),
    ("pop3s.{}", 110, 0),
    ("pop-mail.{}", 995, 1),
    ("pop-mail.{}", 110, 0),
    ("pop3-mail.{}", 995, 1),
    ("pop3-mail.{}", 110, 0),
    ("popmail.{}", 995, 1),
    ("popmail.{}", 110, 0),
    ("secure-pop.{}", 995, 1),
    ("secure-pop.{}", 110, 0),
    ("ssl-pop.{}", 995, 1),
    ("ssl-pop.{}", 110, 0),
    ("securemail.{}", 995, 1),
    ("securemail.{}", 110, 0),
    ("pop1.{}", 995, 1),
    ("pop1.{}", 110, 0),
    ("pop2.{}", 995, 1),
    ("pop2.{}", 110, 0),
    ("pop01.{}", 995, 1),
    ("pop01.{}", 110, 0),
    ("pop02.{}", 995, 1),
    ("pop02.{}", 110, 0),
    ("pop3-01.{}", 995, 1),
    ("pop3-01.{}", 110, 0),
    ("pop3-02.{}", 995, 1),
    ("pop3-02.{}", 110, 0),
    ("mail1.{}", 995, 1),
    ("mail1.{}", 110, 0),
    ("mail2.{}", 995, 1),
    ("mail2.{}", 110, 0),
    ("mail01.{}", 995, 1),
    ("mail01.{}", 110, 0),
    ("mail02.{}", 995, 1),
    ("mail02.{}", 110, 0),
    ("mailhost.{}", 995, 1),
    ("mailhost.{}", 110, 0),
    ("mailserver.{}", 995, 1),
    ("mailserver.{}", 110, 0),
    ("email.{}", 995, 1),
    ("email.{}", 110, 0),
    ("inbox.{}", 995, 1),
    ("inbox.{}", 110, 0),
    ("mailbox.{}", 995, 1),
    ("mailbox.{}", 110, 0),
    ("webmail.{}", 995, 1),
    ("webmail.{}", 110, 0),
    ("incoming.{}", 995, 1),
    ("incoming.{}", 110, 0),
    ("inbound.{}", 995, 1),
    ("inbound.{}", 110, 0),
    ("receive.{}", 995, 1),
    ("receive.{}", 110, 0),
    ("recv.{}", 995, 1),
    ("recv.{}", 110, 0),
    ("mx.{}", 995, 1),
    ("mx.{}", 110, 0),
    ("mx1.{}", 995, 1),
    ("mx1.{}", 110, 0),
    ("mx2.{}", 995, 1),
    ("mx2.{}", 110, 0),
]


# ==========================================================
# UNIQUE OUTPUT FILE
# ==========================================================

def get_unique_output(filename):
    """
    Return an unused filename.

    Example:
        imap.csv
        imap-(2).csv
        imap-(3).csv
        imap-(4).csv
    """

    path = Path(filename)

    # Original filename does not exist
    if not path.exists():
        return str(path)

    counter = 2

    while True:

        new_path = path.with_name(
            f"{path.stem}-({counter}){path.suffix}"
        )

        if not new_path.exists():
            return str(new_path)

        counter += 1


# ==========================================================
# CLEAN DOMAIN
# ==========================================================

def clean_domain(line):

    domain = line.strip().strip('"').strip("'").lower()

    if not domain:
        return None

    domain = domain.replace("https://", "")
    domain = domain.replace("http://", "")

    # Remove possible www.
    if domain.startswith("www."):
        domain = domain[4:]

    domain = domain.split("/")[0]
    domain = domain.split("?")[0]
    domain = domain.split("#")[0]
    domain = domain.rstrip(".")

    if "." not in domain:
        return None

    return domain


# ==========================================================
# PORT CHECK
# ==========================================================

def check_port(domain, template, port, ssl):

    host = template.format(domain)

    try:

        with socket.create_connection(
            (host, port),
            timeout=TIMEOUT
        ):

            return (
                domain,
                host,
                port,
                ssl
            )

    except (
        socket.timeout,
        socket.gaierror,
        ConnectionRefusedError,
        TimeoutError,
        OSError
    ):
        return None


# ==========================================================
# CHECK DOMAIN
# ==========================================================

def check_domain(domain):

    smtp = []
    imap = []
    pop3 = []

    # ======================================================
    # SMTP
    # ======================================================

    for template, port, ssl in SMTP_SERVERS:

        result = check_port(
            domain,
            template,
            port,
            ssl
        )

        if result:
            smtp.append(result)

    # ======================================================
    # IMAP
    # ======================================================

    for template, port, ssl in IMAP_SERVERS:

        result = check_port(
            domain,
            template,
            port,
            ssl
        )

        if result:
            imap.append(result)

    # ======================================================
    # POP3
    # ======================================================

    for template, port, ssl in POP3_SERVERS:

        result = check_port(
            domain,
            template,
            port,
            ssl
        )

        if result:
            pop3.append(result)

    return smtp, imap, pop3


# ==========================================================
# WRITE RESULT
# ==========================================================

def write_result(file_handle, row, seen):

    key = (
        row[0],
        row[1],
        row[2],
        row[3]
    )

    if key in seen:
        return False

    seen.add(key)

    writer = csv.writer(file_handle)
    writer.writerow(row)
    file_handle.flush()

    return True


# ==========================================================
# MAIN
# ==========================================================

def iter_input_lines(path):
    """Yield text values from TXT/LOG/CSV/XLSX input files."""

    suffix = path.suffix.lower()

    if suffix in {".txt", ".log"}:
        with open(path, "r", encoding="utf-8", errors="ignore") as handle:
            yield from handle
        return

    if suffix == ".csv":
        with open(
            path,
            "r",
            encoding="utf-8-sig",
            errors="ignore",
            newline=""
        ) as handle:
            reader = csv.reader(handle)
            for row in reader:
                for value in row:
                    if value and value.strip():
                        yield value
        return

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError(
                "XLSX input requires openpyxl. "
                "Install it with: pip install openpyxl"
            )

        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True
        )

        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    for value in row:
                        if value is not None:
                            yield str(value)
        finally:
            workbook.close()

        return

    raise RuntimeError(f"Unsupported input type: {path.suffix}")


def main():

    # ======================================================
    # INPUT FILE
    # ======================================================

    while True:

        input_file = input(
            "Please input file or folder: "
        ).strip().strip('"')

        if not input_file:
            print(
                "[ERROR] Please enter a file or folder."
            )
            continue

        input_path = Path(input_file)

        # --------------------------------------------------
        # Single file
        # --------------------------------------------------

        if input_path.is_file():

            files = [input_path]
            break

        # --------------------------------------------------
        # Folder
        # --------------------------------------------------

        if input_path.is_dir():

            files = sorted(
                p
                for p in input_path.iterdir()
                if p.is_file()
                and p.suffix.lower() in {
                    ".txt",
                    ".log",
                    ".csv",
                    ".xlsx"
                }
            )

            if files:
                break

            print(
                "[ERROR] No .txt, .log, .csv or .xlsx files found."
            )
            continue

        print(
            f"[ERROR] File or folder not found: {input_path}"
        )


    # ======================================================
    # THREADS
    # ======================================================

    while True:

        try:

            threads = int(
                input(
                    "Please input number of threads: "
                ).strip()
            )

            if threads <= 0:
                raise ValueError

            break

        except ValueError:

            print(
                "[ERROR] Threads must be a positive number."
            )


    # ======================================================
    # UNIQUE OUTPUT FILES
    # ======================================================

    SMTP_OUTPUT_FILE = get_unique_output(
        SMTP_OUTPUT
    )

    IMAP_OUTPUT_FILE = get_unique_output(
        IMAP_OUTPUT
    )

    POP3_OUTPUT_FILE = get_unique_output(
        POP3_OUTPUT
    )


    # ======================================================
    # DISPLAY
    # ======================================================

    print()
    print("=" * 70)
    print("SMTP + IMAP + POP3 SERVER CHECKER")
    print("=" * 70)

    print("Input:")

    for file in files:
        print(f"  - {file}")

    print()
    print(f"Threads : {threads}")
    print(f"Timeout : {TIMEOUT}s")

    print()
    print("Output:")

    print(f"  SMTP : {SMTP_OUTPUT_FILE}")
    print(f"  IMAP : {IMAP_OUTPUT_FILE}")
    print(f"  POP3 : {POP3_OUTPUT_FILE}")

    print("=" * 70)
    print()


    # ======================================================
    # STATISTICS
    # ======================================================

    total_read = 0
    valid_domains = 0
    checked = 0

    smtp_found = 0
    imap_found = 0
    pop3_found = 0

    duplicate_domains = 0


    # ======================================================
    # DOMAIN-LEVEL DEDUPLICATION
    # ======================================================

    seen_domains = set()


    # ======================================================
    # RESULT-LEVEL DEDUPLICATION
    # ======================================================

    seen_smtp = set()
    seen_imap = set()
    seen_pop3 = set()


    # ======================================================
    # OPEN OUTPUT
    # ======================================================

    with open(
        SMTP_OUTPUT_FILE,
        "w",
        encoding="utf-8",
        newline=""
    ) as smtp_file, open(
        IMAP_OUTPUT_FILE,
        "w",
        encoding="utf-8",
        newline=""
    ) as imap_file, open(
        POP3_OUTPUT_FILE,
        "w",
        encoding="utf-8",
        newline=""
    ) as pop3_file:

        # --------------------------------------------------
        # CSV HEADERS
        # --------------------------------------------------

        smtp_file.write(
            "domain,server,port,ssl\n"
        )

        imap_file.write(
            "domain,server,port,ssl\n"
        )

        pop3_file.write(
            "domain,server,port,ssl\n"
        )

        smtp_file.flush()
        imap_file.flush()
        pop3_file.flush()


        # ==================================================
        # THREAD POOL
        # ==================================================

        with ThreadPoolExecutor(
            max_workers=threads
        ) as executor:

            pending = set()


            # ==================================================
            # SAVE COMPLETED RESULT
            # ==================================================

            def save_results(future):

                nonlocal checked
                nonlocal smtp_found
                nonlocal imap_found
                nonlocal pop3_found

                try:

                    smtp, imap, pop3 = future.result()

                except Exception:
                    return

                checked += 1


                # ------------------------------------------
                # SMTP
                # ------------------------------------------

                for row in smtp:

                    if write_result(
                        smtp_file,
                        row,
                        seen_smtp
                    ):

                        smtp_found += 1


                # ------------------------------------------
                # IMAP
                # ------------------------------------------

                for row in imap:

                    if write_result(
                        imap_file,
                        row,
                        seen_imap
                    ):

                        imap_found += 1


                # ------------------------------------------
                # POP3
                # ------------------------------------------

                for row in pop3:

                    if write_result(
                        pop3_file,
                        row,
                        seen_pop3
                    ):

                        pop3_found += 1


                # ------------------------------------------
                # STATUS
                # ------------------------------------------

                print(
                    f"Read:{total_read:,} "
                    f"Checked:{checked:,} "
                    f"SMTP:{smtp_found:,} "
                    f"IMAP:{imap_found:,} "
                    f"POP3:{pop3_found:,} "
                    f"Pending:{len(pending):,}",
                    end="\r",
                    flush=True
                )


            # ==================================================
            # READ ALL FILES
            # ==================================================

            for current_file in files:

                print()
                print(
                    f"[INPUT] {current_file}"
                )

                try:

                    for line in iter_input_lines(current_file):

                            total_read += 1

                            domain = clean_domain(line)

                            if not domain:
                                continue


                            # ----------------------------------
                            # DOMAIN DEDUPLICATION
                            # ----------------------------------

                            if domain in seen_domains:

                                duplicate_domains += 1
                                continue

                            seen_domains.add(domain)

                            valid_domains += 1


                            # ----------------------------------
                            # SUBMIT TASK
                            # ----------------------------------

                            pending.add(
                                executor.submit(
                                    check_domain,
                                    domain
                                )
                            )


                            # ----------------------------------
                            # LIMIT PENDING TASKS
                            # ----------------------------------

                            if len(pending) >= MAX_PENDING:

                                done, pending = wait(
                                    pending,
                                    return_when=FIRST_COMPLETED
                                )

                                for task in done:
                                    save_results(task)


                except Exception as error:

                    print()
                    print(
                        f"[ERROR] Cannot read "
                        f"{current_file}: {error}"
                    )


            # ==================================================
            # FINISH REMAINING TASKS
            # ==================================================

            while pending:

                done, pending = wait(
                    pending,
                    return_when=FIRST_COMPLETED
                )

                for task in done:
                    save_results(task)


    # ==========================================================
    # FINAL OUTPUT
    # ==========================================================

    print()
    print()

    print("=" * 70)
    print("[DONE]")
    print("=" * 70)

    print(
        f"Files processed       : {len(files):,}"
    )

    print(
        f"Lines read            : {total_read:,}"
    )

    print(
        f"Unique domains        : {valid_domains:,}"
    )

    print(
        f"Duplicate domains     : {duplicate_domains:,}"
    )

    print(
        f"Domains checked       : {checked:,}"
    )

    print(
        f"SMTP found            : {smtp_found:,}"
    )

    print(
        f"IMAP found            : {imap_found:,}"
    )

    print(
        f"POP3 found            : {pop3_found:,}"
    )

    print()
    print("Output files:")

    print(
        f"  SMTP : {SMTP_OUTPUT_FILE}"
    )

    print(
        f"  IMAP : {IMAP_OUTPUT_FILE}"
    )

    print(
        f"  POP3 : {POP3_OUTPUT_FILE}"
    )

    print("=" * 70)


# ==========================================================
# START
# ==========================================================

if __name__ == "__main__":
    main()
